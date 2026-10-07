# compare_to_actual.py - compares the generated valuation (after Excel recalculated it) with the analyst's real workbook
# Usage: python compare_to_actual.py --generated "<generated saved by Excel>.xlsx" --outputs <generated>_outputs.json
#          --actual "<actual workbook>.xlsx" --actual-cells actual_cells.json --out-dir /app/created [--tolerance 0.005]
# actual_cells.json: {"enterprise_value": "Valuation!B2", "valuation_ebitda": "Valuation!B1", ...}; metrics allowed:
#   ltm_revenue, valuation_ebitda, ev_ebitda_multiple, enterprise_value, net_debt, equity_value, fair_value, xpv_value
# Explains the EV gap exactly: EBITDA effect + multiple effect (using the actual workbook's implied multiple).
import sys, os, json, argparse
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import parse_ref, is_num

ORDER = ["ltm_revenue", "valuation_ebitda", "ev_ebitda_multiple", "enterprise_value", "net_debt", "equity_value",
         "fair_value", "xpv_value"]
NAMES = {"ltm_revenue": "LTM revenue", "valuation_ebitda": "EBITDA the multiple is applied to",
         "ev_ebitda_multiple": "EV/EBITDA multiple (input)", "enterprise_value": "Enterprise value",
         "net_debt": "Net debt", "equity_value": "Equity value before discount",
         "fair_value": "Equity value after discount", "xpv_value": "XPV value"}


def main(a):
    meta = json.load(open(a.outputs, encoding="utf-8"))
    cells = json.load(open(a.actual_cells, encoding="utf-8"))
    bad = [k for k in cells if k not in ORDER]
    if bad:
        print(f"COMPARE_ERROR: unknown metrics {bad}; allowed {ORDER}"); return 1
    g = load_workbook(a.generated, data_only=True)
    t = load_workbook(a.actual, data_only=True)
    gv = lambda ref: (lambda k: g[k[0]].cell(row=k[1], column=k[2]).value)(parse_ref(ref))
    tv = lambda ref: (lambda k: t[k[0]].cell(row=k[1], column=k[2]).value)(parse_ref(ref))
    G = {k: gv(meta["outputs"][k]) for k in ORDER if k in meta["outputs"]}
    if all(v is None for v in G.values()):
        print("COMPARE STATUS: NOT_RECALCULATED\nNEXT REQUIRED STEP: ask the user to open the generated workbook in desktop "
              "Excel, save it, and upload the saved file."); return 3
    T = {k: tv(ref) for k, ref in cells.items()}
    checks = [(c, gv(c)) for c in meta["checks"]]
    rows, n_cmp, n_ok = [], 0, 0
    for k in ORDER:
        if k not in T:
            continue
        gval, aval = G.get(k), T[k]
        if is_num(gval) and is_num(aval):
            d = gval - aval
            pct = d / abs(aval) if aval else None
            ok = abs(pct) <= a.tolerance if pct is not None else abs(d) < 1e-9
            n_cmp += 1; n_ok += ok
            rows.append(f"| {NAMES[k]} | {gval:,.2f} | {aval:,.2f} ({cells[k]}) | {d:+,.2f} | "
                        f"{'' if pct is None else f'{pct:+.2%}'} | {'MATCH' if ok else 'DIFFERENT'} |")
        else:
            rows.append(f"| {NAMES[k]} | {gval} | {aval} ({cells[k]}) | | | NOT COMPARABLE |")
    bridge = []
    if all(is_num(x) for x in (G.get("enterprise_value"), G.get("valuation_ebitda"), T.get("enterprise_value"),
                                T.get("valuation_ebitda"))) and T["valuation_ebitda"]:
        ge, gm = G["valuation_ebitda"], G["enterprise_value"] / G["valuation_ebitda"] if G["valuation_ebitda"] else 0
        ae, am = T["valuation_ebitda"], T["enterprise_value"] / T["valuation_ebitda"]
        e_eff, m_eff = (ge - ae) * am, (gm - am) * ge
        bridge = ["", "## Enterprise value gap, explained",
                  f"- Implied EV/EBITDA: generated {gm:.4f}, actual {am:.4f}"
                  + (f" (actual input multiple {T['ev_ebitda_multiple']})" if is_num(T.get("ev_ebitda_multiple")) else ""),
                  f"- EBITDA effect: ({ge:,.2f} - {ae:,.2f}) x {am:.4f} = {e_eff:+,.2f}",
                  f"- Multiple effect: ({gm:.4f} - {am:.4f}) x {ge:,.2f} = {m_eff:+,.2f}",
                  f"- Total EV gap: {e_eff + m_eff:+,.2f}"]
        if is_num(T.get("ev_ebitda_multiple")) and abs(am - T["ev_ebitda_multiple"]) > 1e-6:
            bridge.append("- The actual workbook's implied multiple differs from its input multiple: its EV formula "
                          "applies further terms (e.g. a factor, a second method, or other adjustments). Review that formula.")
        if all(is_num(x) for x in (G.get("net_debt"), T.get("net_debt"))):
            bridge.append(f"- Net debt gap: {G['net_debt'] - T['net_debt']:+,.2f} (reduces equity by this amount)")
    failed_checks = [c for c, v in checks if v is not True]
    IMP = {"xpv_value": 3, "fair_value": 3, "equity_value": 2.8, "enterprise_value": 2.8, "valuation_ebitda": 2.5,
           "net_debt": 2.2, "ltm_revenue": 2.0, "ev_ebitda_multiple": 2.0}
    review = []
    for k in ORDER:
        if k in T and is_num(G.get(k)) and is_num(T[k]) and T[k]:
            pct = abs((G[k] - T[k]) / T[k])
            if pct > a.tolerance:
                review.append((round(IMP[k] * min(1.0, 0.3 + pct / 0.1), 2), NAMES[k], f"generated differs by {pct:.1%}"))
    for c in failed_checks:
        review.append((3.0, c, "generated workbook check fails"))
    for x in meta.get("assumptions", []):
        review.append((1.2, "assumption", x))
    review.sort(key=lambda r: -r[0])
    status = ("GENERATED_MATCHES" if n_cmp and n_ok == n_cmp and not failed_checks else "GENERATED_DIFFERS")
    L = [f"# Generated vs actual: {meta['workbook']} vs {os.path.basename(a.actual)}", "", f"**{status}** "
         f"({n_ok}/{n_cmp} metrics within {a.tolerance:.1%})", "",
         f"EBITDA route: {' + '.join(meta['ebitda_route'])} | bridge: {meta['bridge_route']} | period end: {meta['period_end']}", "",
         "| Metric | Generated | Actual | Difference | % | Result |", "|---|---|---|---|---|---|"] + rows + bridge
    L[4:4] = ["## Review list (highest priority first: impact x uncertainty)", "| # | Priority | Item | Why |", "|---|---|---|---|"] + \
             [f"| {i} | {r[0]:.2f} | {r[1]} | {r[2]} |" for i, r in enumerate(review, 1)] + [""]
    L += ["", "## Generated workbook checks"] + [f"- {'PASS' if v is True else 'FAIL'}: {c}" for c, v in checks]
    L += ["", "_Differences show where this company's workbook uses company-specific logic that the generic template "
          "does not have. Each one is a candidate for an analyst answer or a company-specific rule._"]
    os.makedirs(a.out_dir, exist_ok=True)
    rep = os.path.join(a.out_dir, os.path.splitext(meta["workbook"])[0] + "_vs_actual.md")
    open(rep, "w", encoding="utf-8").write("\n".join(L))
    print(f"COMPARE STATUS: {status} | {n_ok}/{n_cmp} metrics within {a.tolerance:.1%} | failed checks: {len(failed_checks)}")
    for line in rows + bridge:
        print(line)
    print(f"REPORT WRITTEN: {rep}")
    print("NEXT REQUIRED STEP: run write_manifest.py, save the files, and give the final summary: COMPARE STATUS line, "
          "the metrics table and the EV gap explanation unchanged, then one line per DIFFERENT metric naming the analyst "
          "question or company-specific rule it points to.")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--generated", required=True); ap.add_argument("--outputs", required=True)
    ap.add_argument("--actual", required=True); ap.add_argument("--actual-cells", required=True)
    ap.add_argument("--out-dir", default="/app/created"); ap.add_argument("--tolerance", type=float, default=0.005)
    try:
        sys.exit(main(ap.parse_args()))
    except Exception as e:
        print(f"COMPARE_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
