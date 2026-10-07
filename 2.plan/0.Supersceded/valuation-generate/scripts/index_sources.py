# index_sources.py (v3) - indexes every number in the source statements with its meaning, and checks the pack
# Usage: python index_sources.py <files or folders> --out /app/created/sources_index.json [--expect-layout <mapping_profile.json>]
# Understands: one file with months across columns; one file per month (current / YTD / budget / variance
# columns); one sheet per month (month in the sheet name); two-row headers (year above month); CSV.
# Repeated labels on one sheet (e.g. consolidated and subsidiary blocks) are kept apart as "Label #2", "#3".
# Writes a layout check: layout type per file, and date problems (duplicate months, gaps, undated files,
# future dates). Problems never stop the run; they are listed for the end-of-run review.
import os, re, sys, csv, json, glob, argparse, datetime, hashlib, bisect
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as col_let
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import parse_period, parse_date_text, column_role, ym_to_i, ym_str, norm_label

UPLOAD_SUFFIX = re.compile(r"-[0-9a-f]{8}(?=\.[A-Za-z0-9]+$)")
MONTH_ONLY = re.compile(r"^(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?$")
_MON = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
HEADING = re.compile(r"\b(as at|as of|ended|ending|period|month|quarter|year)\b")


def stable_name(path):
    return UPLOAD_SUFFIX.sub("", os.path.basename(path))


def grids(path):
    if path.lower().endswith(".csv"):
        g = {}
        with open(path, newline="", encoding="utf-8-sig") as fh:
            for r, row in enumerate(csv.reader(fh), 1):
                for c, v in enumerate(row, 1):
                    v = v.strip()
                    if not v:
                        continue
                    try:
                        g[(r, c)] = float(v.replace(",", "").replace("$", "").replace("(", "-").replace(")", ""))
                    except ValueError:
                        g[(r, c)] = v
        yield "csv", g
        return
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        g = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and getattr(c, "row", None):
                    g[(c.row, c.column)] = c.value
        yield ws.title, g


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def txt(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime.datetime) else str(v).strip()


def sheet_values(fn, sheet, g):
    """numbers with label, header, role, header period; fast nearest-text lookups."""
    text_by_col, text_by_row = defaultdict(list), defaultdict(list)
    for (r, c), v in g.items():
        if not is_num(v) and v is not None and txt(v):
            text_by_col[c].append(r)
            if isinstance(v, str) and not parse_period(v.strip()):
                text_by_row[r].append(c)
    for k in text_by_col:
        text_by_col[k].sort()
    for k in text_by_row:
        text_by_row[k].sort()
    out, occ = [], defaultdict(int)
    for (r, c) in sorted(k for k, v in g.items() if is_num(v)):
        v = g[(r, c)]
        cols = text_by_row.get(r, [])
        i = bisect.bisect_left(cols, c)
        label = str(g[(r, cols[i - 1])]).strip()[:80] if i else "?"
        rows = text_by_col.get(c, [])
        j = bisect.bisect_left(rows, r)
        parts = [txt(g[(rr, c)]) for rr in rows[max(0, j - 3):j]]
        h = " | ".join(parts)
        role = column_role(h)
        per = None
        for part in parts[::-1]:
            d = parse_period(part) or (parse_date_text(part) if role in ("current", "prior_year") else None)
            if d:
                per = d
                break
        if per is None and len(parts) > 1 and role in ("current", "prior_year"):
            per = parse_date_text(" ".join(parts))                      # year and month on two header rows
        bare = bool(parts) and MONTH_ONLY.match(parts[-1].lower().strip())
        if per is None and bare:                                        # "Aug" with the year written once, to the left
            hr = rows[j - 1]
            years = [(cc, int(m_.group(0))) for (rr, cc), vv in g.items() if rr < hr and cc <= c
                     and isinstance(vv, (str, int, float)) and not isinstance(vv, bool)
                     for m_ in [re.fullmatch(r"(19|20)\d{2}", str(vv).strip().split(".")[0])] if m_]
            if years:
                yr = max(years)[1]
                per = (yr, _MON[parts[-1].lower().strip()[:3]])
        undated_month = bare and per is None
        occ[(norm_label(label), c)] += 1
        n = occ[(norm_label(label), c)]
        out.append({"file": fn, "sheet": sheet, "cell": f"{col_let(c)}{r}", "row": r, "col": c, "label": label,
                    "label_key": label if n == 1 else f"{label} #{n}", "header": h, "role": role,
                    "period": f"{per[0]}-{per[1]:02d}" if per and role != "variance" else "", "value": float(v),
                    "undated_month_header": undated_month})
    return out


def main(paths, out, next_step, expect):
    files = []
    for p in paths:
        files += ([f for ext in ("xlsx", "xlsm", "csv") for f in glob.glob(os.path.join(p, "**", "*." + ext), recursive=True)]
                  if os.path.isdir(p) else [p])
    files = sorted(f for f in files if not os.path.basename(f).startswith("~$"))
    if not files:
        raise FileNotFoundError("no .xlsx/.xlsm/.csv source files found")
    names = [stable_name(f) for f in files]
    if len(set(names)) != len(names):
        raise ValueError("two source files share a name once upload suffixes are removed: " + ", ".join(names))
    values, finfo, checks = [], [], []
    for f in files:
        fn = stable_name(f)
        sheets = list(grids(f))
        fp = parse_date_text(os.path.splitext(fn)[0]); fp_src = "file name" if fp else None
        if not fp:
            for _, g in sheets:
                hit = next((parse_date_text(v) for (r, c), v in sorted(g.items())
                            if r <= 8 and isinstance(v, str) and HEADING.search(v.lower()) and parse_date_text(v)), None)
                if hit:
                    fp, fp_src = hit, "sheet heading"; break
        sheet_dates = {sh: parse_date_text(sh) for sh, _ in sheets}
        dated_sheets = {sh: d for sh, d in sheet_dates.items() if d}
        mine = []
        for sh, g in sheets:
            vals = sheet_values(fn, sh, g)
            sd = dated_sheets.get(sh) if len(dated_sheets) > 1 else None
            if sd:                                          # one sheet per month: same statement, many months
                base = re.sub(r"[\s_\-]*(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s_\-]*\d*.*$", "", sh.lower())
                for x in vals:
                    x["sheet_key"] = base.strip() or "statement"
                    if not x["period"] and x["role"] in ("current", "ytd", "budget"):
                        x["period"] = f"{sd[0]}-{sd[1]:02d}"
            mine += vals
        header_dates = sorted({x["period"] for x in mine if x["period"]})
        if not fp and len(dated_sheets) > 1:
            fp, fp_src = max(dated_sheets.values(), key=ym_to_i), "latest dated sheet"
        if not fp and header_dates:
            fp, fp_src = parse_period(header_dates[-1]), "latest dated column"
        undated = [x for x in mine if x.get("undated_month_header")]
        if undated:
            checks.append(f"{fn}: {len(undated)} numbers are under month headers with no year; they are not used")
        if fp:
            for x in mine:
                if x.get("undated_month_header"):
                    continue
                if not x["period"] and x["role"] in ("current", "ytd", "budget"):
                    x["period"] = f"{fp[0]}-{fp[1]:02d}"
                elif not x["period"] and x["role"] == "prior_year":
                    x["period"] = f"{fp[0] - 1}-{fp[1]:02d}"
        months = {x["period"] for x in mine if x["period"] and x["role"] == "current"}
        layout = ("sheet per month" if len(dated_sheets) > 1 else
                  "months across columns" if len(months) >= 3 else
                  "one period per file" if fp else "unknown dates")
        dups = sum(1 for x in mine if x["label_key"] != x["label"])
        values += mine
        finfo.append({"name": fn, "uploaded_as": os.path.basename(f), "sha256": hashlib.sha256(open(f, "rb").read()).hexdigest(),
                      "statement_period": f"{fp[0]}-{fp[1]:02d}" if fp else None, "statement_period_source": fp_src,
                      "layout": layout, "sheets": [sh for sh, _ in sheets], "numbers": len(mine),
                      "roles": sorted({x["role"] for x in mine}), "repeated_label_rows": dups})
        if not fp and not months:
            checks.append(f"{fn}: no date found (file name, heading, sheet names or column headers) - its numbers cannot be used")
        if dups:
            checks.append(f"{fn}: {dups} numbers sit under a label repeated on the same sheet; they are kept apart as '<label> #2' etc.")
    # pack-level date sanity
    single = [fi for fi in finfo if fi["layout"] == "one period per file" and fi["statement_period"]]
    by_p = defaultdict(list)
    for fi in single:
        by_p[fi["statement_period"]].append(fi["name"])
    for p_, ns in by_p.items():
        if len(ns) > 1:
            checks.append(f"{len(ns)} files claim the same month {p_}: {', '.join(ns)} (the latest-named is used where values differ)")
    if len(by_p) >= 2:
        idxs = sorted(ym_to_i(parse_period(p_)) for p_ in by_p)
        gaps = [ym_str(i) for i in range(idxs[0], idxs[-1]) if i not in idxs]
        if gaps:
            checks.append(f"missing monthly statements for {', '.join(gaps)}")
    today = datetime.date.today()
    for fi in finfo:
        if fi["statement_period"] and ym_to_i(parse_period(fi["statement_period"])) > today.year * 12 + today.month:
            checks.append(f"{fi['name']}: statement period {fi['statement_period']} is in the future - check the file date")
    layout_now = sorted({(fi["layout"], tuple(fi["roles"])) for fi in finfo})
    if expect:
        prof = json.load(open(expect, encoding="utf-8"))
        before = sorted({(l["layout"], tuple(l["roles"])) for l in prof.get("statement_layout", [])})
        if before and before != layout_now:
            checks.append(f"statement layout changed since the profile was built: was {before}, now {layout_now}")
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    json.dump({"version": 3, "files": finfo, "layout_checks": checks, "values": values}, open(out, "w", encoding="utf-8"))
    roles = defaultdict(int)
    for v in values:
        roles[v["role"]] += 1
    print(f"SOURCES INDEXED: {len(files)} files | numbers: {len(values)} | column roles: "
          + ", ".join(f"{k}={n}" for k, n in sorted(roles.items())))
    for fi in finfo:
        print(f"- {fi['name']}: {fi['layout']}, statement period {fi['statement_period'] or 'UNKNOWN'}"
              + (f" (from {fi['statement_period_source']})" if fi["statement_period"] else ""))
    print("LAYOUT CHECK: " + ("OK" if not checks else f"{len(checks)} item(s) for the end-of-run review"))
    for c in checks:
        print(f"- {c}")
    print(f"FILE WRITTEN: {out}")
    print(f"NEXT REQUIRED STEP: {next_step} Do not stop for layout-check items; list them in the final summary.")


if __name__ == "__main__":
    NEXT_STEP = "run map_financials.py."
    ap = argparse.ArgumentParser()
    ap.add_argument("paths", nargs="+"); ap.add_argument("--out", default="/app/created/sources_index.json")
    ap.add_argument("--expect-layout", help="mapping_profile.json whose statement_layout the new pack should match")
    a = ap.parse_args()
    try:
        main(a.paths, a.out, NEXT_STEP, a.expect_layout)
    except Exception as e:
        print(f"INDEX_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
