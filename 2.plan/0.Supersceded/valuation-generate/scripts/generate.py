# generate.py - baseline valuation for a company WITHOUT a workbook: one command per job (deterministic)
#   build    --statements S [S ...] --company "Display Name" [--answers /app/created/answers.json]
#   compare  --saved F [--actual A]        (after the Excel step; --actual = the analyst's real workbook, optional)
# Global: --out (deliverables, default /app/created)  --work (scratch, default /tmp/valuation_work)  --debug
# Exit codes: 0 done | 10 judgment needed | 20 Excel step | 3 cannot build | 1 error
# Per run: <company_id>_<quarter>_generated.xlsx and <company_id>_<quarter>_generate_log.json (mapping, assumptions, results).
import os, re, sys, json, glob, shutil, argparse, subprocess, datetime
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
LIMIT, UPLOADS = 4_000_000, "/app/uploads"
# files can be uploaded to SharePoint by path only from here (the self-test points it at its own folder)
UPLOAD_ROOT = os.environ.get("VALUATION_UPLOAD_ROOT") or ("/app/created" if os.path.isdir("/app") else None)
UPLOAD_SUFFIX = re.compile(r"-[0-9a-f]{8}(?=\.[A-Za-z0-9]+$)|(?:_\d{1,2}| \(\d{1,2}\))(?=\.[A-Za-z0-9]+$)")
SUFFIX = {"inc", "ltd", "limited", "llc", "lp", "llp", "corp", "corporation", "co", "company", "technologies", "technology",
          "holdings", "holding", "group", "plc", "gmbh", "ag", "sa", "bv", "ulc"}
METRICS = [("ltm_revenue", r"(ltm|ttm|trailing|last twelve).*(revenue|sales)|(revenue|sales).*\b(ltm|ttm)\b"),
           ("valuation_ebitda", r"ebitda(?!.*(margin|%|multiple|/))"),
           ("ev_ebitda_multiple", r"(\bev\b|enterprise value)\s*/\s*ebitda|ebitda multiple|\bmultiple\b"),
           ("enterprise_value", r"enterprise value|^\s*ev\b(?!\s*/)"),
           ("net_debt", r"net debt"), ("equity_value", r"equity value"), ("fair_value", r"fair value"),
           ("xpv_value", r"\bxpv\b")]
SKIP = re.compile(r"prior|previous|budget|variance|growth|margin avg|check", re.I)


def sh(args):
    p = subprocess.run([sys.executable] + args, capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def line(out, key):
    return next((l.split(":", 1)[1].strip() for l in out.splitlines() if l.startswith(key)), "")


def load(p, default=None):
    return json.load(open(p, encoding="utf-8")) if p and os.path.exists(p) else default


def dump(obj, p):
    json.dump(obj, open(p, "w", encoding="utf-8"), indent=1, default=str)
    return p


def company_slug(name):
    words = re.findall(r"[A-Za-z0-9]+", name or "")
    core = [w for w in words if w.lower() not in SUFFIX] or words
    return "-".join(w.lower() for w in core)[:40] or None


def quarter_of(p):
    return f"{p[:4]}-Q{(int(p[5:7]) - 1) // 3 + 1}" if p and re.match(r"\d{4}-\d{2}", str(p)) else "period"


def resolve(path, out):
    if not path:
        return None
    if os.path.exists(path):
        return os.path.abspath(path)
    want = UPLOAD_SUFFIX.sub("", os.path.basename(path))
    for d in (out, UPLOADS):
        for f in glob.glob(os.path.join(d, "**", "*"), recursive=True) if os.path.isdir(d) else []:
            if os.path.isfile(f) and UPLOAD_SUFFIX.sub("", os.path.basename(f)) == want:
                return f
    return None


def sp_lines(cid, sub, paths):
    rel = f"Valuations/{cid}/{sub}"
    L = [f"SHAREPOINT: Create new folder | library: Documents | path: {rel}  (if it already exists, continue)"]
    for p_ in paths:
        if not p_ or not os.path.exists(p_):
            continue
        n = os.path.basename(p_)
        if os.path.getsize(p_) >= LIMIT:
            L.append(f"SHAREPOINT: SKIP {n} (>= 4 MB) - return it as a download")
        elif UPLOAD_ROOT and not os.path.abspath(p_).startswith(UPLOAD_ROOT.rstrip("/") + "/"):
            L.append(f"SHAREPOINT: SKIP {n} - not in {UPLOAD_ROOT}; return it as a download")
        else:
            L.append(f"SHAREPOINT: Create file | folderPath: /Shared Documents/{rel} | name: {n} | body: {os.path.abspath(p_)}")
    return L


def work(a, cid):
    d = os.path.join(a.out, "debug", cid, "generate") if a.debug else os.path.join(a.work, cid, "generate")
    os.makedirs(d, exist_ok=True)
    return d


def build(a):
    out = a.out; os.makedirs(out, exist_ok=True)
    cid = company_slug(a.company)
    if not cid:
        print("PIPELINE ERROR: give --company \"<display name>\""); sys.exit(1)
    W = work(a, cid)
    idx = os.path.join(W, "sources_index.json")
    rc, o = sh([os.path.join(HERE, "index_sources.py")] + a.statements + ["--out", idx])
    if "INDEX_ERROR" in o:
        print("PIPELINE ERROR in index statements:\n" + "\n".join(l for l in o.splitlines() if "ERROR" in l)); sys.exit(1)
    a.statements = [resolve(x, out) or x for x in a.statements]
    missing = [x for x in a.statements if not os.path.exists(x)]
    if missing:
        print("PIPELINE ERROR: these files are not in the chat: " + "; ".join(missing) +
              ". Ask the user to upload them or paste their SharePoint links (download each with sharepoint_get_doc)."); sys.exit(1)
    ans = None
    if a.answers:                                   # one-line JSON, '-' for stdin (heredoc), or a path
        raw = sys.stdin.read() if a.answers.strip() == "-" else a.answers
        if raw.strip().startswith("{"):
            ans = dump(json.loads(raw), os.path.join(W, "answers.json"))
        else:
            ans = resolve(raw, out)
    mp = os.path.join(W, "mapping.json")
    rc, o = sh([os.path.join(HERE, "map_financials.py"), "--sources-index", idx, "--out", mp] + (["--answers", ans] if ans else []))
    ms = line(o, "MAPPING STATUS")
    used = [l for l in o.splitlines() if l.startswith("- ") and not l.endswith(": missing")]
    missing = [l[2:].split(":")[0] for l in o.splitlines() if l.startswith("- ") and l.endswith(": missing")]
    flags = [l for l in o.splitlines() if l.startswith(("FLAG:", "ASSUMPTION:", "BLOCKER:"))]
    print(f"MAPPING STATUS: {ms}")
    if "MAPPING_ERROR" in o or ms.startswith("CANNOT_BUILD"):
        print("\n".join(l for l in o.splitlines() if l.startswith(("BLOCKER:", "MAPPING_ERROR"))))
        print("NEXT: stop. Post the lines above as one list."); sys.exit(3)
    q = quarter_of((load(mp) or {}).get("period_end"))
    wb = os.path.join(out, f"{cid}_{q}_generated.xlsx")
    b_args = [os.path.join(HERE, "build_workbook.py"), "--mapping", mp, "--out", wb] + (["--answers", ans] if ans else [])
    rc, o2 = sh(b_args)
    if not line(o2, "BUILD STATUS").startswith("BUILT"):
        print("\n".join(l for l in o2.splitlines() if l.startswith(("BUILD", "- "))) or o2[-800:]); sys.exit(3)
    op = wb.rsplit(".", 1)[0] + "_outputs.json"
    meta = load(op); os.replace(op, os.path.join(W, os.path.basename(op)))
    log = dump({"company_id": cid, "display_name": a.company, "quarter": q, "period_end": meta.get("period_end"),
                "status": "awaiting_excel", "created_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
                "mapping_status": ms, "statement_lines_used": [l[2:] for l in used], "not_found": missing,
                "flags_and_assumptions": flags, "answers": load(ans) if ans else None, "outputs": meta,
                "inputs": {"statements": [os.path.basename(s) for s in a.statements]},
                "paths": {"workbook": os.path.abspath(wb)}}, os.path.join(out, f"{cid}_{q}_generate_log.json"))
    print(f"BUILD STATUS: {line(o2, 'BUILD STATUS').split('|')[0].strip()} | {a.company} (company_id {cid}) {q}")
    print("Statement lines used:")
    for l in used:
        print(l)
    if missing:
        print(f"Not found in the statements: {', '.join(missing)}")
    print(f"Review ({len(flags)}):" + ("" if flags else " nothing needs review."))
    for i, f in enumerate(flags, 1):
        print(f"{i}. {f}")
    print(f"EXCEL STEP: return {os.path.basename(wb)} with references/excel-step.md. When the saved file is uploaded run: "
          f"python scripts/generate.py compare --saved <uploaded file> [--actual <analyst's workbook>]")
    print(f"FILES: {os.path.basename(wb)}, {os.path.basename(log)} (needed only to continue in a new chat)")
    print("NEXT: give the user the workbook with the Excel-step instructions and stop.")
    sys.exit(20)


def candidates(path):
    """closed list of cells in the analyst's workbook that could be each metric (label to the left of a value)."""
    from openpyxl import load_workbook
    wf, wv = load_workbook(path, data_only=False), load_workbook(path, data_only=True)
    found = {m: [] for m, _ in METRICS}
    for ws in wf.worksheets:
        pref = 0 if re.search(r"summary|valuation|output|report", ws.title, re.I) else 1
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or c.value.startswith("=") or SKIP.search(c.value):
                    continue
                vals = [x for x in row if x.column > c.column and x.column <= c.column + 8 and
                        (isinstance(x.value, (int, float)) and not isinstance(x.value, bool)
                         or (isinstance(x.value, str) and x.value.startswith("=")))]
                if not vals:
                    continue
                v = vals[-1] if len(vals) >= 3 else vals[0]
                for m, rx in METRICS:
                    if re.search(rx, c.value, re.I):
                        cached = wv[ws.title][v.coordinate].value
                        found[m].append((pref, f"'{ws.title}'!{v.coordinate}",
                                         f"'{ws.title}'!{v.coordinate} label '{c.value[:50]}' {str(v.value)[:60]} -> {cached}"
                                         + (f" (last of {len(vals)} columns)" if len(vals) >= 3 else "")))
    qs = []
    for m, _ in METRICS:
        opts = sorted(found[m], key=lambda x: x[0])[:4]
        if opts:
            qs.append({"metric": m, "options": {chr(65 + i): o[2] for i, o in enumerate(opts)},
                       "_cells": {chr(65 + i): o[1] for i, o in enumerate(opts)}})
    return qs


def compare(a):
    out = a.out
    logs = [p for d in (out, UPLOADS) if os.path.isdir(d) for p in glob.glob(os.path.join(d, "**", "*_generate_log*.json"), recursive=True)]
    lp = resolve(a.log, out) if a.log else (max(logs, key=os.path.getmtime) if logs else None)
    saved = resolve(a.saved, out)
    if not lp or not saved:
        print("PIPELINE ERROR: upload " + ", ".join(x for x, ok in (("the saved generated workbook", saved),
              ("the <company>_<quarter>_generate_log.json from the build step", lp)) if not ok)); sys.exit(1)
    log = load(lp); cid, q = log["company_id"], log["quarter"]
    W = work(a, cid)
    op = dump(log["outputs"], os.path.join(W, "outputs.json"))
    if not a.actual:
        rc, o = sh([os.path.join(HERE, "read_values.py"), "--generated", saved, "--outputs", op])
        st = line(o, "VALUES STATUS")
        print(f"VALUES STATUS: {st} | {log['display_name']} ({cid}) {q}")
        if st != "READ":
            print("NEXT: ask the user to open the workbook in desktop Excel, let it calculate, save and upload it. Stop."); sys.exit(0)
        vals = [l for l in o.splitlines() if not l.startswith(("VALUES STATUS", "NEXT"))]
        print("\n".join(vals))
        log.update(status="values_read", values=vals)
    else:
        actual = resolve(a.actual, out)
        if not actual:
            print(f"PIPELINE ERROR: actual workbook not found: {a.actual}"); sys.exit(1)
        allq = candidates(actual)
        qs = [x for x in allq if len(x["options"]) > 1]            # a single candidate is chosen by rule, listed to confirm
        pk = os.path.join(out, "actual_cell_picks.json")
        if qs and not os.path.exists(pk):
            dump({"task": "choose the actual workbook's cells", "write": pk,
                  "format": {"<metric>": "<letter>", "...": "or \"none\" when no option clearly is that metric"},
                  "rules": "Choose from the label and formula shown; skip rather than guess. Do not open the workbook.",
                  "questions": [{k: v for k, v in x.items() if k != "_cells"} for x in qs]}, os.path.join(out, "judgment.json"))
            print(f"JUDGMENT NEEDED: {len(qs)} question(s) in {os.path.join(out, 'judgment.json')}. Write {pk}, then re-run the same command.")
            sys.exit(10)
        picks = load(pk) or {}
        pick = lambda x: "A" if len(x["options"]) == 1 else str(picks.get(x["metric"], "none")).strip().upper()[:1]
        cells = {x["metric"]: x["_cells"][pick(x)] for x in allq if pick(x) in x["_cells"]}
        by_ai = {x["metric"] for x in qs}
        ac = dump({k: v.replace("'", "") if re.match(r"^'[^' ]+'!", v) else v for k, v in cells.items()}, os.path.join(W, "actual_cells.json"))
        rc, o = sh([os.path.join(HERE, "compare_to_actual.py"), "--generated", saved, "--outputs", op, "--actual", actual,
                    "--actual-cells", ac, "--out-dir", W])
        st = line(o, "COMPARE STATUS")
        print(f"COMPARE STATUS: {st} | {log['display_name']} ({cid}) {q}")
        if "NOT_RECALCULATED" in st:
            print("NEXT: ask the user to open the workbook in desktop Excel, let it calculate, save and upload it. Stop."); sys.exit(0)
        if "COMPARE_ERROR" in o:
            print(o[-800:]); sys.exit(1)
        rows = [l for l in o.splitlines() if l.startswith(("|", "- ", "##"))]
        print("| Metric | Generated | Actual | Difference | % | Result |\n|---|---|---|---|---|---|")
        print("\n".join(rows))
        print("Actual cells used (confirm): " + ", ".join(f"{k} {v}{' (AI pick)' if k in by_ai else ''}" for k, v in cells.items()))
        log.update(status=st, actual=os.path.basename(actual), actual_cells=cells, comparison=rows)
        if os.path.exists(pk):
            os.replace(pk, os.path.join(W, "actual_cell_picks.json"))
        if os.path.exists(os.path.join(out, "judgment.json")):
            os.replace(os.path.join(out, "judgment.json"), os.path.join(W, "judgment.json"))
    if log.get("flags_and_assumptions"):
        print("Assumptions and flags from the build: " + "; ".join(log["flags_and_assumptions"]))
    dump(log, lp)
    final = os.path.join(out, f"{cid}_{q}_generated.xlsx")
    if os.path.abspath(saved) != os.path.abspath(final):
        shutil.copy(saved, final)
    for l in sp_lines(cid, q, [final, lp]):
        print(l)
    print("NEXT: carry out the SHAREPOINT lines, then post this report unchanged. Stop.")
    sys.exit(0)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    for n in ("build", "compare"):
        p = sub.add_parser(n)
        p.add_argument("--out", default="/app/created" if os.path.isdir("/app") else os.getcwd())
        p.add_argument("--work", default="/tmp/valuation_work"); p.add_argument("--debug", action="store_true")
    b = sub.choices["build"]
    b.add_argument("--statements", nargs="+", required=True); b.add_argument("--company", required=True); b.add_argument("--answers")
    c = sub.choices["compare"]
    c.add_argument("--saved", required=True); c.add_argument("--actual"); c.add_argument("--log")
    a = ap.parse_args()
    {"build": build, "compare": compare}[a.cmd](a)
