# read_values.py - prints the generated valuation's recalculated values and checks (no actual workbook)
# Usage: python read_values.py --generated "<generated saved by Excel>.xlsx" --outputs <generated>_outputs.json
import sys, os, json, argparse
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import parse_ref


def main(a):
    meta = json.load(open(a.outputs, encoding="utf-8"))
    g = load_workbook(a.generated, data_only=True)
    v = lambda ref: (lambda k: g[k[0]].cell(row=k[1], column=k[2]).value)(parse_ref(ref))
    vals = {k: v(r) for k, r in meta["outputs"].items()}
    if all(x is None for x in vals.values()):
        print("VALUES STATUS: NOT_RECALCULATED\nNEXT REQUIRED STEP: ask the user to open the workbook in desktop Excel, save, "
              "and upload it."); return 3
    print("VALUES STATUS: READ")
    for k, x in vals.items():
        print(f"{k} ({meta['outputs'][k]}): {x}")
    for c in meta["checks"]:
        print(f"check {'PASS' if v(c) is True else 'FAIL'}: {c}")
    print("NEXT REQUIRED STEP: give the final summary with these values exactly as printed.")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--generated", required=True); ap.add_argument("--outputs", required=True)
    sys.exit(main(ap.parse_args()))
