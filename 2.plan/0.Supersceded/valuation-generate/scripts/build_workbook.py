# build_workbook.py - builds a generic valuation workbook from mapping.json + answers.json
# Usage: python build_workbook.py --mapping mapping.json --answers answers.json --out <company>_valuation_<period>.xlsx
# Python copies source values and writes FORMULAS only; Excel does every calculation.
# Also writes <out>_outputs.json (metric -> cell) for compare_to_actual.py.
import sys, os, json, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

UNITS = {"units": 1, "ones": 1, "dollars": 1, "thousands": 1000, "000s": 1000, "k": 1000, "millions": 1000000, "m": 1000000}
BLUE, BOLD = Font(color="0000FF"), Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")


def main(a):
    m = json.load(open(a.mapping, encoding="utf-8"))
    ans = json.load(open(a.answers, encoding="utf-8"))
    if m["status"] not in ("READY", "READY_WITH_ASSUMPTIONS"):
        print("BUILD STATUS: REFUSED\n- mapping status is " + m["status"] + "; report its blockers in the final summary")
        return 1
    ans.setdefault("source_units", "units"); ans.setdefault("report_units", "thousands")
    for f in ("source_units", "report_units"):
        if str(ans.get(f, "")).lower() not in UNITS:
            print(f"BUILD STATUS: REFUSED\n- {f} must be one of {sorted(UNITS)}"); return 1
    wb = Workbook()
    inp = wb.active; inp.title = "Inputs"
    inp.append(["Input", "Value", "Source"]); [setattr(c, "font", BOLD) for c in inp[1]]
    rows = [("company", "Company", ans.get("company")), ("period_end", "Period end", m["period_end"]),
            ("src_mult", "Source units multiplier", UNITS[str(ans["source_units"]).lower()]),
            ("rep_mult", "Report units multiplier", UNITS[str(ans["report_units"]).lower()]),
            ("factor", "Units factor (source to report)", "=B4/B5"),
            ("m_ebitda", "EV/EBITDA multiple", ans.get("ev_ebitda_multiple")),
            ("m_rev", "EV/Revenue multiple", ans.get("ev_revenue_multiple") or 0),
            ("w_ebitda", "Weight EV/EBITDA", ans.get("weight_ev_ebitda", 1)),
            ("w_rev", "Weight EV/Revenue", ans.get("weight_ev_revenue", 0)),
            ("adj", "EBITDA adjustments (report units)", ans.get("ebitda_adjustments", 0)),
            ("disc", "Illiquidity discount", ans.get("illiquidity_discount", 0)),
            ("own", "XPV fully diluted ownership", ans.get("ownership_pct", 1)),
            ("claims", "Claims ahead of common equity (report units)", ans.get("other_claims", 0))]
    I = {}
    for i, (k, lab, v) in enumerate(rows, start=2):
        inp.cell(i, 1, lab); c = inp.cell(i, 2, v)
        c.font = Font(color="000000") if k == "factor" else BLUE
        inp.cell(i, 3, "formula" if k == "factor" else "analyst (answers.json)")
        I[k] = f"Inputs!$B${i}"
    inp.column_dimensions["A"].width = 44; inp.column_dimensions["B"].width = 16; inp.column_dimensions["C"].width = 24

    src = wb.create_sheet("Source Data")
    src.append(["Concept", "Source line", "File", "Sheet", "Row", "Frequency", "Values (as in source) ->"])
    for c in src[1]:
        c.font, c.fill = BOLD, HEAD
    used = set(m["ebitda_route"]) | {"revenue"} | ({"net_debt"} if m["bridge_route"] == "net_debt" else {"debt", "cash"})
    R, lineage = {}, []
    r = 2
    for cid in [c for c in m["concepts"] if c in used and m["concepts"][c]["status"] == "mapped"]:
        con = m["concepts"][cid]
        ranges = []
        for line in con["lines"]:
            src.cell(r, 1, cid); src.cell(r, 2, line["label"]); src.cell(r, 3, line["file"])
            src.cell(r, 4, line["sheet"]); src.cell(r, 5, line["row"]); src.cell(r, 6, line["frequency"])
            for j, (p, v) in enumerate(zip(line["periods"], line["values"])):
                src.cell(1, 8 + j, p) if not src.cell(1, 8 + j).value else None
                src.cell(r, 8 + j, v)
            rng = f"'Source Data'!$H${r}:${L(7 + len(line['values']))}${r}"
            ranges.append(rng)
            lineage.append([cid, line["label"], line["file"], line["sheet"], line["row"], line["frequency"],
                            f"{line['periods'][0]}..{line['periods'][-1]}", con["match"], rng])
            r += 1
        R[cid] = ranges
    src.column_dimensions["B"].width = 30; src.column_dimensions["C"].width = 28

    def amt(cid):
        s = "+".join(f"SUM({x})" for x in R[cid])
        s = f"ABS({s})" if m["concepts"][cid]["cost"] else f"({s})"
        return f"{s}*{I['factor']}"

    val = wb.create_sheet("Valuation")
    val.append(["Line", "Value", "Formula basis"]); [setattr(c, "font", BOLD) for c in val[1]]
    V, vr = {}, 2

    def put(key, label, formula, basis):
        nonlocal vr
        val.cell(vr, 1, label); val.cell(vr, 2, formula); val.cell(vr, 3, basis)
        V[key] = f"Valuation!$B${vr}"; vr += 1

    route = m["ebitda_route"]
    put("ltm_revenue", "LTM revenue", "=" + amt("revenue") if "revenue" in R else 0, "sum of the LTM window x units factor")
    if route == ["adjusted_ebitda"]:
        put("ltm_ebitda", "LTM EBITDA (reported adjusted)", "=" + amt("adjusted_ebitda"), "statement line")
    elif route == ["ebitda"]:
        put("ltm_ebitda", "LTM EBITDA (reported)", "=" + amt("ebitda"), "statement line")
    elif route == ["ebit", "da"]:
        put("ltm_ebitda", "LTM EBITDA (EBIT + D&A)", f"={amt('ebit')}+{amt('da')}", "derived: EBIT + D&A")
    else:
        put("ltm_ebitda", "LTM EBITDA (revenue - COGS - opex)", f"={amt('revenue')}-{amt('cogs')}-{amt('opex')}",
            "derived: revenue - COGS - operating expenses")
    put("valuation_ebitda", "EBITDA for valuation (incl. adjustments)", f"={V['ltm_ebitda']}+{I['adj']}", "LTM EBITDA + analyst adjustments")
    put("ev_ebitda", "EV from EV/EBITDA", f"={V['valuation_ebitda']}*{I['m_ebitda']}", "EBITDA x multiple")
    put("ev_revenue", "EV from EV/Revenue", f"={V['ltm_revenue']}*{I['m_rev']}", "revenue x multiple")
    put("enterprise_value", "Enterprise value", f"={V['ev_ebitda']}*{I['w_ebitda']}+{V['ev_revenue']}*{I['w_rev']}", "weighted")
    if m["bridge_route"] == "net_debt":
        put("net_debt", "Net debt", "=" + amt("net_debt"), "statement line at period end")
    else:
        put("net_debt", "Net debt (debt - cash)", f"={amt('debt')}-{amt('cash')}", "derived at period end")
    put("claims", "Claims ahead of common equity", f"={I['claims']}", "analyst")
    put("equity_value", "Equity value (before discount)", f"={V['enterprise_value']}-{V['net_debt']}-{V['claims']}", "EV - net debt - claims")
    put("fair_value", "Equity value after discount", f"={V['equity_value']}*(1-{I['disc']})", "x (1 - discount)")
    put("xpv_value", "XPV value", f"={V['fair_value']}*{I['own']}", "x ownership")
    val.column_dimensions["A"].width = 42; val.column_dimensions["B"].width = 16; val.column_dimensions["C"].width = 40

    chk = wb.create_sheet("Checks")
    chk.append(["Check", "Result (TRUE = pass)"]); [setattr(c, "font", BOLD) for c in chk[1]]
    checks = [("Method weights sum to 1", f"=ABS({I['w_ebitda']}+{I['w_rev']}-1)<0.000001")]
    for cid, rngs in R.items():
        if m["concepts"][cid]["kind"] == "flow":
            n = {"monthly": 12, "quarterly": 4, "annual": 1}[m["concepts"][cid]["lines"][0]["frequency"]]
            checks.append((f"{cid}: LTM window has {n} values", f"=COUNT({rngs[0]})={n}"))
    if all(c in R for c in ("net_debt", "debt", "cash")):
        checks.append(("Net debt line ties to debt - cash (informational)",
                       f"=ABS({amt('net_debt')}-({amt('debt')}-{amt('cash')}))<1"))
    C = []
    for i, (lab, f) in enumerate(checks, start=2):
        chk.cell(i, 1, lab); chk.cell(i, 2, f); C.append(f"Checks!$B${i}")
    chk.column_dimensions["A"].width = 52

    lin = wb.create_sheet("Lineage")
    lin.append(["Concept", "Source line", "File", "Sheet", "Row", "Frequency", "Periods", "Match", "Source Data range"])
    for c in lin[1]:
        c.font, c.fill = BOLD, HEAD
    for row in lineage:
        lin.append(row)
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    wb.save(a.out)
    outs = {k: v.replace("$", "") for k, v in V.items()}
    outs["ev_ebitda_multiple"] = I["m_ebitda"].replace("$", "")
    meta = {"workbook": os.path.basename(a.out), "outputs": outs, "checks": [c.replace("$", "") for c in C],
            "assumptions": m.get("assumptions", []), "layout_checks": m.get("layout_checks", []),
            "ebitda_route": route, "bridge_route": m["bridge_route"], "period_end": m["period_end"]}
    op = a.out.rsplit(".", 1)[0] + "_outputs.json"
    json.dump(meta, open(op, "w", encoding="utf-8"), indent=1)
    print(f"BUILD STATUS: BUILT | {a.out} | EBITDA route: {' + '.join(route)} | bridge: {m['bridge_route']} | checks: {len(C)}")
    print(f"OUTPUT MAP: {op}")
    print("NEXT REQUIRED STEP: run write_manifest.py, then give the user the workbook with the Excel-step instructions "
          "(references/excel-step.md) and end your turn. When the saved file comes back, run compare_to_actual.py if an "
          "actual workbook was provided, otherwise report the generated values with read_values.py.")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--mapping", required=True); ap.add_argument("--answers", required=True); ap.add_argument("--out", required=True)
    try:
        sys.exit(main(ap.parse_args()))
    except Exception as e:
        print(f"BUILD_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
