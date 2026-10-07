# check_result.py - checks a draft after Excel recalculated it (deterministic)
# Usage:
#   python check_result.py --mode replay|backtest|refresh --base <base.xlsx> [--actual <real new-period workbook> (backtest)]
#          [--draft <draft.xlsx>] --input-map <input_map.json> --plan <run_log.json or write log>
#          (--result <saved_after_excel.xlsx> | --values-json <office_script_output.json>)
#          --out-json <check.json> [--report-md <check_report.md> (debug only)]
#          [--input-threshold 0.25] [--output-threshold 0.15]
# Checks: 1) inputs written as planned  2) formulas unchanged vs base  3) Excel recalculated
#         4) no new Excel errors (#REF!, #VALUE!, #DIV/0!, #N/A, #NAME?, #NUM!, #NULL!)
#         5) replay: key outputs equal the base; backtest: inputs and key outputs equal the actual workbook;
#            refresh: variance flags vs last quarter  6) the workbook's own controls pass
# Nothing here recalculates the valuation: Excel is the model; this only checks that the update landed cleanly.
import sys, json, argparse, os
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import formulas_of, norm_formula, parse_ref, ref, is_num, cells_of

ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def close(a, b):
    if is_num(a) and is_num(b):
        return abs(a - b) <= max(1e-6, 1e-7 * max(abs(a), abs(b)))
    return a == b


def error_cells(book):
    out = {}
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRORS:
                    out[f"{ws.title}!{c.coordinate}"] = c.value.strip()
    return out


def main(a):
    imap = json.load(open(a.input_map, encoding="utf-8"))
    plan = json.load(open(a.plan, encoding="utf-8"))
    base_f = load_workbook(a.base, data_only=False)
    base_v = load_workbook(a.base, data_only=True)
    truth_v = load_workbook(a.actual, data_only=True) if a.mode == "backtest" else base_v
    res_f = load_workbook(a.result, data_only=False) if a.result else None
    draft_f = load_workbook(a.draft, data_only=False) if a.draft and os.path.exists(a.draft) else res_f
    writes = plan["writes"]

    # 1) inputs written as planned
    bad_in = []
    AUTO = ("replace_period_value", "roll_window")
    cur_not_rebuilt = [p["cell"] for p in writes if p["write_mode"] in AUTO and p.get("offset_months") == 0
                       and (p.get("origin") or {}).get("type") != "rule"]
    historical = [p for p in writes if p["status"] == "historical"]
    for p in writes:
        if not p.get("written") or p["status"] == "historical":
            continue
        sh, cell = p["cell"].rsplit("!", 1)
        for label, book in (("draft", draft_f), ("saved file", res_f)):
            got = book[sh][cell].value if book is not None else p["value"]
            if not close(got, p["value"]):
                bad_in.append(f"{p['cell']}: planned {p['value']}, {label} has {got}")
                break
        if a.mode == "replay" and not close(p["value"], p["base_value"]):
            bad_in.append(f"{p['cell']}: rebuilt {p['value']} != original {p['base_value']} [{p['concept_id']}]")
        if a.mode == "backtest" and not close(p["value"], p.get("actual_value")):
            bad_in.append(f"{p['cell']}: predicted {p['value']} != actual {p.get('actual_value')} [{p['concept_id']}]")
    exceptions = [p for p in writes if p["status"] == "exception"]

    # 2) formulas unchanged
    fb = formulas_of(base_f)
    changed, added = set(), set()
    for book in [b for b in (draft_f, res_f) if b is not None]:
        fd = formulas_of(book)
        changed |= {ref(k) for k in fb if k not in fd or norm_formula(fb[k]) != norm_formula(fd[k])}
        added |= {ref(k) for k in fd if k not in fb}
    changed, added = sorted(changed), sorted(added)

    # 3) recalculated values for key outputs and controls
    kos = imap.get("key_outputs", [])
    ctrls = imap.get("controls", [])
    cells = sorted(set(kos + [c["cell"] for c in ctrls]))
    res_v, script_errors = None, None
    if a.values_json:
        got = json.load(open(a.values_json, encoding="utf-8"))
        if isinstance(got, dict) and "values" in got:          # Office Script recalc-and-read format
            script_errors = {}
            for e in got.get("errors") or []:
                cell, _, err = str(e).rpartition("=")
                script_errors[cell.replace("'", "")] = err
            got = got["values"]
        got = {k.replace("'", ""): v for k, v in got.items()}
        rv = lambda c: got.get(c.replace("'", ""))
    else:
        res_v = load_workbook(a.result, data_only=True)
        def rv(c):
            k = parse_ref(c)
            return res_v[k[0]].cell(row=k[1], column=k[2]).value
    bv = lambda c: (lambda k: truth_v[k[0]].cell(row=k[1], column=k[2]).value)(parse_ref(c))
    not_calc = [c for c in cells if rv(c) is None and bv(c) is not None]

    # 4) new Excel errors (errors already in the base workbook are not new)
    fprec = {k: set(v) for k, v in (imap.get("key_output_formula_precedents") or {}).items()}
    feeding = set(cells).union(*fprec.values()) if fprec else set(cells)
    new_err = {}
    if res_v is not None or script_errors is not None:
        before = {k.replace("'", "") for k in error_cells(base_v)}
        after = error_cells(res_v) if res_v is not None else script_errors
        new_err = {c: e for c, e in after.items() if c.replace("'", "") not in before}
    feeding = {x.replace("'", "") for x in feeding}
    new_err = {c.replace("'", ""): e for c, e in new_err.items()}
    err_block = sorted(c for c in new_err if c in feeding)
    err_other = sorted(c for c in new_err if c not in feeding)

    # 5) outputs
    ko_rows, out_bad, var_flags, var_meta = [], [], [], []
    for c in kos:
        o, n = bv(c), rv(c)
        if a.mode in ("replay", "backtest"):
            ok = close(o, n)
            ko_rows.append({"cell": c, "before": o, "after": n, "result": "MATCH" if ok else "DIFFERENT"})
            if not ok:
                out_bad.append(c)
        else:
            chg = (n - o) / abs(o) if is_num(o) and is_num(n) and o != 0 else None
            flag = chg is not None and abs(chg) > a.output_threshold
            ko_rows.append({"cell": c, "before": o, "after": n, "change": None if chg is None else round(chg, 4),
                            "flag": bool(flag)})
            if flag:
                var_flags.append(f"key output {c} moved {chg:+.1%}")
                var_meta.append((c, None, chg, a.output_threshold, True))
    if a.mode == "refresh":
        for p in writes:
            o, n = p["base_value"], p["value"]
            if p.get("written") and is_num(o) and is_num(n) and o != 0 and abs((n - o) / abs(o)) > a.input_threshold:
                var_flags.append(f"input {p['cell']} [{p['concept_id']}] moved {(n - o) / abs(o):+.1%}")
                var_meta.append((p["cell"], p["concept_id"], (n - o) / abs(o), a.input_threshold, False))

    # 6) controls (the workbook's own checks)
    c_rows, c_block = [], []
    for c in ctrls:
        v = rv(c["cell"])
        if isinstance(v, str) and v.startswith("#"):
            res = "FAIL"
        elif c["control_type"] == "boolean":
            res = "PASS" if isinstance(v, bool) and v == bool(c["expected"]) else "FAIL"
        else:
            res = "PASS" if is_num(v) and abs(v - c["expected"]) <= float(c.get("absolute_tolerance", 0)) else "FAIL"
        c_rows.append({"cell": c["cell"], "value": v, "expected": c["expected"], "result": res, "severity": c["severity"]})
        if res == "FAIL" and c["severity"] == "blocking":
            c_block.append(c["cell"])

    in_label = {"replay": " and equal to the original", "backtest": " and equal to the actual workbook", "refresh": ""}[a.mode]
    out_label = {"replay": "key outputs reproduce the original", "backtest": "key outputs match the actual workbook",
                 "refresh": "no variance flags"}[a.mode]
    checks = {"every current-period input rebuilt by its rule": not cur_not_rebuilt,
              "inputs written as planned" + in_label: not bad_in,
              "no planned cell left as an exception": not exceptions,
              "formulas unchanged": not changed and not added,
              "Excel recalculated the file": not not_calc,
              "no new Excel errors in key outputs or their precedents": not err_block,
              out_label: not out_bad if a.mode in ("replay", "backtest") else not var_flags,
              "blocking controls pass": not c_block}
    if a.mode in ("replay", "backtest"):
        status = f"{a.mode.upper()}_PASS" if all(checks.values()) else f"{a.mode.upper()}_FAIL"
    else:
        hard = [k for k in checks if k != "no variance flags" and not checks[k]]
        status = "DRAFT_BLOCKED" if hard else "DRAFT_READY"

    # ---- why does each differing key output differ? (deterministic attribution to precedents)
    feeds = {}
    for e in imap.get("inputs", []):
        for k in cells_of(e["id"]):
            feeds[ref(k)] = set(e.get("feeds_key_outputs") or [])
    target_of = lambda p: p.get("actual_value") if a.mode == "backtest" else p["base_value"]
    changed_in = [p for p in writes if p.get("written") and not close(p["value"], target_of(p))]
    explain = {}
    for c in out_bad:
        ins = [p for p in changed_in if c in feeds.get(p["cell"], set())]
        fs_ = [x for x in changed + added if x in fprec.get(c, set())]
        explain[c] = (ins, fs_)
    unexplained = [c for c, (ins, fs_) in explain.items() if not ins and not fs_]
    fx_irrelevant = [x for x in changed + added if not any(x in v for v in fprec.values())] if fprec else []

    # ---- one review list, highest priority (impact x uncertainty) first; empty when nothing needs a person
    review = []
    def item(impact, unc, what, why, cell=None, concept=None):
        review.append({"priority": round(impact * unc, 2), "item": what, "why": why, "cell": cell, "concept": concept})
    for p in exceptions:
        item(3, 1.0, f"{p['cell']} [{p['concept_id']}]", f"not written: {p['message']}", p["cell"], p["concept_id"])
    for c in c_block:
        item(3, 1.0, c, "blocking control fails", c)
    for c in err_block:
        item(3, 1.0, c, f"new Excel error {new_err[c]} (feeds a key output)", c)
    for c in cur_not_rebuilt:
        item(3, 1.0, c, "current-period input not rebuilt by its rule", c)
    for x in bad_in:
        item(3, 0.9, x.split(":")[0], x, x.split(":")[0])
    for c in out_bad:
        ins, fs_ = explain[c]
        why = ("UNEXPLAINED: no written input or changed formula feeds it - check external links, volatile functions "
               "(TODAY, NOW, OFFSET) and the recalculation") if c in unexplained else (
               "explained by " + "; ".join([f"{p['cell']} [{p['concept_id']}] {target_of(p)} -> {p['value']}" for p in ins[:4]]
                                           + [f"formula {x} changed" for x in fs_[:3]])
               + (f" (+{len(ins) + len(fs_) - 7} more)" if len(ins) + len(fs_) > 7 else ""))
        item(3, 1.0, f"key output {c}", why, c)
    for x in [x for x in changed + added if x not in fx_irrelevant]:
        item(3, 0.9, x, "formula changed or added, and it feeds a key output", x)
    if fx_irrelevant:
        item(1, 0.4, f"{len(fx_irrelevant)} changed formulas ({fx_irrelevant[0]} ...)",
             "formula changed or added, but it feeds no key output (display or notes); confirm")
    if err_other:
        item(1, 0.5, f"{len(err_other)} new Excel errors outside the valuation path ({err_other[0]} {new_err[err_other[0]]} ...)",
             "display or notes cells; confirm they are harmless")
    for c in not_calc:
        item(3, 0.8, c, "not recalculated by Excel", c)
    for f_, (cell, cid, pct, thr, is_ko) in zip(var_flags, var_meta):
        item(3 if is_ko else 2, min(1.0, 0.4 + 0.3 * abs(pct) / max(thr, 1e-9)),
             f"{'key output ' if is_ko else ''}{cell}" + (f" [{cid}]" if cid else ""),
             f"moved {pct:+.1%} vs last quarter (threshold {thr:.0%})", cell, cid)
    groups = {}
    for p in writes:
        if p["status"] in ("flag", "historical") and not (a.mode == "replay" and p["status"] == "historical"):
            groups.setdefault((p["status"], p["concept_id"], p["message"]), []).append(p["cell"])
    for (st_, cid, msg), cells_ in groups.items():
        what = f"{cells_[0]} [{cid}]" if len(cells_) == 1 else f"{cid}: {len(cells_)} cells ({cells_[0]} ... {cells_[-1]})"
        item(2 if st_ == "flag" else 1.5, 0.5 if st_ == "flag" else 0.3, what, msg, cells_[0], cid)
    for n in plan.get("notes", []):
        item(1.5 if "layout differs" in n or "replay only" in n else 1, 0.4, "note", n)
    review.sort(key=lambda r: -r["priority"])
    if a.mode == "refresh" and status == "DRAFT_READY" and review:
        status = "DRAFT_READY_FOR_REVIEW"

    result = {"status": status, "checks": checks, "review": review, "key_outputs": ko_rows, "controls": c_rows,
              "new_errors": new_err, "historical_cells_kept": len(historical),
              "counts": {"cells": len(writes), "written": sum(1 for p in writes if p.get("written")),
                         "key_outputs": len(kos), "key_outputs_ok": len(kos) - len(out_bad) if a.mode != "refresh" else None,
                         "controls": len(ctrls), "controls_pass": sum(1 for r in c_rows if r["result"] == "PASS")}}
    json.dump(result, open(a.out_json, "w", encoding="utf-8"), indent=1, default=str)
    if a.report_md:                                   # debug only
        L = [f"# {a.mode.title()} check: {plan['workbook']}", "", f"**{status}**", "", "## Review list"]
        L += [f"{i}. {r['item']} - {r['why']}" for i, r in enumerate(review, 1)] or ["(nothing to review)"]
        L += ["", "## Checks"] + [f"- {'PASS' if ok else 'FAIL'}: {k}" for k, ok in checks.items()]
        L += ["", "## Key outputs"] + [f"- {json.dumps(r, default=str)}" for r in ko_rows]
        L += ["", "## Controls"] + [f"- {json.dumps(r, default=str)}" for r in c_rows]
        open(a.report_md, "w", encoding="utf-8").write("\n".join(L))
    print(f"CHECK STATUS: {status}")
    for k, ok in checks.items():
        if not ok:
            print(f"check FAIL: {k}")
    print(f"REVIEW ITEMS: {len(review)}")
    print(f"RESULT WRITTEN: {a.out_json}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--mode", choices=["replay", "backtest", "refresh"], required=True)
    ap.add_argument("--actual", help="backtest: the analyst's real workbook for the new period")
    ap.add_argument("--base", required=True); ap.add_argument("--draft")
    ap.add_argument("--input-map", required=True); ap.add_argument("--plan", required=True)
    ap.add_argument("--result"); ap.add_argument("--values-json")
    ap.add_argument("--out-json", required=True); ap.add_argument("--report-md")
    ap.add_argument("--input-threshold", type=float, default=0.25)
    ap.add_argument("--output-threshold", type=float, default=0.15)
    a = ap.parse_args()
    if a.mode == "backtest" and not a.actual:
        print("CHECK_ERROR: backtest needs --actual"); sys.exit(1)
    if not (a.result or a.values_json):
        print("CHECK_ERROR: give --result (workbook saved by Excel) or --values-json (Office Script output)"); sys.exit(1)
    try:
        main(a)
    except Exception as e:
        print(f"CHECK_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
