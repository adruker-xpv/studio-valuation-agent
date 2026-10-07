# wbutil.py - shared helpers for valuation-refresh (same fingerprint algorithm as extract_schema.py v4)
import re, hashlib
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx, get_column_letter as col_let, range_boundaries

STR_LIT = re.compile(r'"[^"]*"')
REF = re.compile(
    r"(?<![A-Za-z0-9_.])"
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!)?"
    r"(\$?)([A-Z]{1,3})(\$?)([0-9]+)"
    r"(?::(\$?)([A-Z]{1,3})(\$?)([0-9]+))?"
    r"(?![A-Za-z0-9_(])")
KEY_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^!,']+))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")


def rel_key(formula, row, col):
    def one(ca, c, ra, r):
        ci, ri = col_idx(c), int(r)
        return (f"R{ri}" if ra else f"R[{ri-row}]") + (f"C{ci}" if ca else f"C[{ci-col}]")

    def sub(m):
        s = m.group(1) or m.group(2) or ""
        s = s + "!" if s else ""
        s += one(m.group(3), m.group(4), m.group(5), m.group(6))
        if m.group(8):
            s += ":" + one(m.group(7), m.group(8), m.group(9), m.group(10))
        return s
    return REF.sub(sub, STR_LIT.sub('""', formula))


def formulas_of(wb):
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                if cell.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                    t = str(getattr(v, "text", v))
                    out[(ws.title, cell.row, cell.column)] = t if t.startswith("=") else "=" + t
    return out


def fingerprint(path):
    wb = load_workbook(path, data_only=False)
    f = formulas_of(wb)
    src = "|".join(wb.sheetnames) + "||" + "|".join(sorted(
        f"{k[0]}!{rel_key(t, k[1], k[2])}@{k[1]},{k[2]}" for k, t in f.items()))
    return hashlib.sha256(src.encode("utf-8")).hexdigest()[:16]


def cells_of(block_id):
    """'Sheet!B2:E2' -> [('Sheet', 2, 2), ...] in row-major, left-to-right order."""
    sheet, rng = block_id.rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
    return [(sheet, r, c) for r in range(r1, (r2 or r1) + 1) for c in range(c1, (c2 or c1) + 1)]


def ref(k):
    return f"{k[0]}!{col_let(k[2])}{k[1]}"


def parse_ref(s):
    m = KEY_REF.match(s)
    sh = (m.group(1) or m.group(2)).strip().replace("''", "'")
    c1, r1, _, _ = range_boundaries(m.group(3).replace("$", ""))
    return (sh, r1, c1)


def norm_formula(t):
    return re.sub(r"\s+", "", str(t)).upper()


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ---- period parsing (source column headers) ----
import datetime as _dt
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _yr(y):
    y = int(y)
    return y + 2000 if y < 100 else y


def parse_period(h):
    """Header -> (year, month) of the period END, or None. Handles ISO dates, 'Jun-26', 'June 2026',
    '2026-06', '06/2026', 'Q2-26', 'Q2 2026', 'FY2026' (-> Dec). Totals, YTD, budget -> None."""
    if h is None:
        return None
    s = str(h).strip().lower()
    if not s or any(w in s for w in ("ytd", "total", "budget", "forecast", "plan", "var", "%")):
        return None
    m = re.match(r"^(\d{4})-(\d{2})(?:-(\d{2}))?", s)
    if m:
        return (int(m.group(1)), int(m.group(2)))
    m = re.match(r"^([a-z]{3})[a-z]*[\s\-_']*(\d{2}|\d{4})$", s)
    if m and m.group(1) in _MONTHS:
        return (_yr(m.group(2)), _MONTHS[m.group(1)])
    m = re.match(r"^(\d{1,2})[/\-](\d{4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return (int(m.group(2)), int(m.group(1)))
    m = re.match(r"^q([1-4])[\s\-_']*(?:fy)?(\d{2}|\d{4})$", s) or re.match(r"^(?:fy)?(\d{2}|\d{4})[\s\-_']*q([1-4])$", s)
    if m:
        q, y = (m.group(1), m.group(2)) if s.startswith("q") else (m.group(2), m.group(1))
        return (_yr(y), int(q) * 3)
    m = re.match(r"^fy[\s\-_']*(\d{2}|\d{4})$", s)
    if m:
        return (_yr(m.group(1)), 12)
    return None


def month_index(ym):
    return ym[0] * 12 + ym[1] - 1


def parse_period_end(s):
    """'2026-09-30' or '2026-09' -> (2026, 9)."""
    m = re.match(r"^(\d{4})-(\d{2})", str(s))
    if not m:
        raise ValueError(f"period end must look like YYYY-MM-DD, got {s!r}")
    return (int(m.group(1)), int(m.group(2)))


# =====================================================================================
# v2: statement semantics (file period, column roles), series across files, operations
# =====================================================================================
_MONTH_WORDS = "january|february|march|april|may|june|july|august|september|october|november|december|" \
               "jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec"


def parse_date_text(s):
    """Free text -> (year, month) or None. Handles 'June 30, 2026', '30 June 2026', 'Jun 30 2026',
    'June_30__2026' (file names), '2026-06-30', '20260630', '202606', 'Jun-26', 'Q2 2026'."""
    if s is None:
        return None
    t = re.sub(r"[_\.]+", " ", str(s).lower())
    m = re.search(r"(20\d{2})[-/ ]?(0[1-9]|1[0-2])(?:[-/ ]?(0[1-9]|[12]\d|3[01]))?(?!\d)", t)
    if m:
        return (int(m.group(1)), int(m.group(2)))
    m = re.search(rf"\b({_MONTH_WORDS})\b[\s,\-]*(?:\d{{1,2}}(?:st|nd|rd|th)?[\s,\-]+)?(20\d{{2}}|\d{{2}})\b", t)
    if m:
        return (_yr(m.group(2)), _MONTHS[m.group(1)[:3]])
    m = re.search(rf"\b\d{{1,2}}(?:st|nd|rd|th)?[\s\-]+({_MONTH_WORDS})\b[\s,\-]+(20\d{{2}}|\d{{2}})\b", t)
    if m:
        return (_yr(m.group(2)), _MONTHS[m.group(1)[:3]])
    m = re.search(rf"\b(20\d{{2}})[\s,\-]+({_MONTH_WORDS})\b", t)      # "2026 | Jun" (two-row headers)
    if m:
        return (int(m.group(1)), _MONTHS[m.group(2)[:3]])
    m = re.search(r"\bq([1-4])[\s\-]*(20\d{2}|\d{2})\b", t)
    if m:
        return (_yr(m.group(2)), int(m.group(1)) * 3)
    return None


def column_role(header_text):
    h = " " + re.sub(r"[^a-z0-9%]+", " ", str(header_text).lower()) + " "
    if re.search(r" (ytd|year to date|cumulative|fytd) ", h):
        return "ytd"
    if re.search(r" (budget|bud|plan|forecast|fcst|target) ", h):
        return "budget"
    if re.search(r" (var|variance|diff|difference|change) |%", h):
        return "variance"
    if re.search(r" (prior year|last year|py|ly) ", h):
        return "prior_year"
    return "current"


def norm_label(s):
    """lowercase, '&' -> 'and', punctuation removed, simple plurals removed ('Revenues' == 'Revenue');
    a trailing ' #N' occurrence marker (repeated labels on one sheet) is kept."""
    s = str(s)
    occ = re.search(r"\s#(\d+)$", s)
    s = s[:occ.start()] if occ else s
    s = s.lower().replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", " ", s)
    words = [w[:-1] if len(w) > 3 and w.endswith("s") and not w.endswith("ss") else w for w in s.split()]
    return " ".join(words) + (f" #{occ.group(1)}" if occ else "")


def ym_to_i(ym):
    return ym[0] * 12 + ym[1] - 1


def i_to_ym(i):
    return (i // 12, i % 12 + 1)


def ym_str(i):
    y, m = i_to_ym(i)
    return f"{y}-{m:02d}"


def build_series(values, file_periods):
    """-> {(sheet_n, label_n, role): {month_index: entry}}. When two files give the same month, the file
    whose statement period IS that month wins, then the latest file (restated figures)."""
    series = {}
    for v in values:
        if not v.get("period"):
            continue
        p = parse_period(v["period"])
        if not p:
            continue
        i = ym_to_i(p)
        key = (norm_label(v["sheet_key"] if v.get("sheet_key") else v["sheet"]),
               norm_label(v.get("label_key") or v["label"]), v.get("role", "current"))
        s = series.setdefault(key, {})
        old = s.get(i)
        if old is None:
            s[i] = v
            continue
        fp_new, fp_old = file_periods.get(v["file"]), file_periods.get(old["file"])
        rank = lambda fp: (fp == i, fp if fp is not None else -1)
        if rank(fp_new) > rank(fp_old):
            s[i] = v
    return series


def find_series(series, src):
    """src = {'sheet': ..., 'label': ..., 'role': ...}; sheet may be None (any sheet)."""
    ln, role = norm_label(src["label"]), src.get("role", "current")
    if src.get("sheet"):
        s = series.get((norm_label(src["sheet"]), ln, role))
        if s is not None:
            return s
    # the sheet is missing or renamed (e.g. a new statement layout): follow the line if exactly one sheet has it
    hits = [v for k, v in series.items() if k[1] == ln and k[2] == role]
    return hits[0] if len(hits) == 1 else None


OPS = ("value_at", "sum_months", "ytd_difference", "components")


def evaluate(spec, series, target_i):
    """Executes one operation for one target period. Returns (value, operands, None) or (None, [], reason).
    Operations:
      value_at        the source value for the period (balances, reported figures, YTD values)
      sum_months      sum of `months` consecutive monthly values ending at the period
      ytd_difference  YTD(period) - YTD(previous quarter end); first fiscal quarter = YTD(period).
                      If YTD(previous quarter end) is not supplied, uses YTD(first month) - month(first month).
      components      sum of sub-specs, each with its own sign (e.g. debt lines, revenue - costs)"""
    op = spec["operation"]
    sc, sg = float(spec.get("scale", 1)), float(spec.get("sign", 1))
    if op == "components":
        tot, ops_ = 0.0, []
        for comp in spec["components"]:
            v, o, why = evaluate(comp, series, target_i)
            if v is None:
                return None, [], why
            tot += v; ops_ += o
        return tot * sc * sg, ops_, None
    s = find_series(series, spec["source"])
    if s is None:
        return None, [], f"no source series '{spec['source'].get('label')}' ({spec['source'].get('role', 'current')})"
    if op == "value_at":
        e = s.get(target_i)
        if e is None:
            return None, [], f"no value for {ym_str(target_i)}"
        return e["value"] * sc * sg, [e], None
    if op == "sum_months":
        n = int(spec.get("months", 3))
        es = [s.get(target_i - k) for k in range(n)]
        if any(e is None for e in es):
            miss = [ym_str(target_i - k) for k in range(n) if s.get(target_i - k) is None]
            return None, [], f"missing months {', '.join(miss)}"
        return sum(e["value"] for e in es) * sc * sg, list(reversed(es)), None
    if op == "ytd_difference":
        fs = int(spec.get("fiscal_year_start_month", 1))
        y, m = i_to_ym(target_i)
        cur = s.get(target_i)
        if cur is None:
            return None, [], f"no YTD value for {ym_str(target_i)}"
        months_into_fy = (m - fs) % 12 + 1
        if months_into_fy <= 3:
            return cur["value"] * sc * sg, [cur], None
        prev = s.get(target_i - 3)
        if prev is not None:
            return (cur["value"] - prev["value"]) * sc * sg, [prev, cur], None
        msrc = dict(spec["source"], role="current")
        ms = find_series(series, msrc)
        first_i = target_i - 2
        a_, b_ = s.get(first_i), (ms or {}).get(first_i)
        if a_ is not None and b_ is not None:
            return (cur["value"] - (a_["value"] - b_["value"])) * sc * sg, [a_, b_, cur], None
        return None, [], f"no YTD for {ym_str(target_i - 3)} and no way to derive it"
    return None, [], f"unknown operation {op}"


def infer_fiscal_start(series):
    """Month where a YTD series equals the same line's monthly value (the first month of the fiscal year)."""
    votes = {}
    for (sh, lb, role), s in series.items():
        if role != "ytd":
            continue
        cur = series.get((sh, lb, "current"), {})
        for i, e in s.items():
            c = cur.get(i)
            if c is not None and abs(e["value"]) > 0 and abs(e["value"] - c["value"]) < 1e-6 * max(1, abs(e["value"])):
                votes[i_to_ym(i)[1]] = votes.get(i_to_ym(i)[1], 0) + 1
    return max(votes, key=votes.get) if votes else None
