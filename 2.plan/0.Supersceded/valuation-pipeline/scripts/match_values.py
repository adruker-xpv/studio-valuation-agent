# match_values.py (v2) - discovers EXECUTABLE rules by reproduction
# Usage: python match_values.py <workbook.xlsx> <input_map.json> <sources_index.json> --out matches.json
#                               [--period-end YYYY-MM-DD] [--fiscal-year-start-month N]
# For every in-scope input block, tries each source series (sheet + label + column role, across all files)
# with each operation (value_at, sum_months 3/12, ytd_difference), unit scale and sign, at each period
# offset (0, -3, -6, -9, -12 months from the period end). A rule that reproduces the workbook's values is
# evidence AND the execution spec the refresh engine will run. Cells it cannot reproduce are reported.
import sys, os, json, bisect, argparse
from collections import Counter, defaultdict
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter as col_let
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import (build_series, evaluate, infer_fiscal_start, parse_period, parse_period_end, ym_to_i, ym_str,
                    norm_label)

SCALES = [1, 0.001, 1000, 0.000001, 1000000]
OFFSETS = [0, -3, -6, -9, -12]
MIN_ABS = 100
OP_RANK = {"value_at": 0, "sum_months": 1, "ytd_difference": 2}


def candidate_specs(series, fs):
    for (sh, lb, role), s in series.items():
        if role in ("budget", "variance"):
            continue
        src = {"sheet": sh, "label": lb, "role": role}
        yield {"operation": "value_at", "source": src}
        if role == "current":
            yield {"operation": "sum_months", "months": 3, "source": src}
            yield {"operation": "sum_months", "months": 12, "source": src}
        if role == "ytd":
            yield {"operation": "ytd_difference", "fiscal_year_start_month": fs, "source": src}


def sig(spec, sc, sg):
    return json.dumps({**spec, "scale": sc, "sign": sg}, sort_keys=True)


def main(a):
    m = json.load(open(a.input_map, encoding="utf-8"))
    idx = json.load(open(a.sources_index, encoding="utf-8"))
    fperiods = {f["name"]: ym_to_i(parse_period(f["statement_period"])) for f in idx.get("files", [])
                if isinstance(f, dict) and f.get("statement_period")}
    series = build_series(idx["values"], fperiods)
    fs = a.fiscal_year_start_month or infer_fiscal_start(series) or 1
    if a.period_end:
        pe = ym_to_i(parse_period_end(a.period_end)); pe_src = "given"
    else:
        cands = list(fperiods.values()) or [i for s in series.values() for i in s]
        pe = max(cands); pe_src = "latest statement period"
    wb = load_workbook(a.workbook, data_only=True)

    # pool of every rule's value at every offset
    pool = []
    for spec in candidate_specs(series, fs):
        for off in OFFSETS:
            v, ops_, why = evaluate(dict(spec, scale=1, sign=1), series, pe + off)
            if v is not None:
                pool.append((v, spec, off, ops_))
    pool.sort(key=lambda x: x[0])
    keys = [p[0] for p in pool]

    def hits(val):
        # whole numbers may be rounded (e.g. thousands) - but only large ones; small whole numbers (share
        # counts, months, percentages) must match exactly, or near-misses become false rules
        tol = 0.51 if float(val).is_integer() and abs(val) >= 1000 else max(abs(val) * 1e-7, 0.005)
        out = []
        for sc in SCALES:
            for sg in (1, -1):
                t = val / (sc * sg)
                lo, hi = bisect.bisect_left(keys, t - tol / sc), bisect.bisect_right(keys, t + tol / sc)
                for v, spec, off, ops_ in pool[lo:hi]:
                    out.append((sig(spec, sc, sg), spec, sc, sg, off, ops_))
        return out

    blocks, tally = [], Counter()
    for e in m["inputs"]:
        if not ((e.get("material") and e.get("type") != "not_material") or e.get("reporting_required") == "yes"):
            continue
        c1, r1, c2, r2 = range_boundaries(e["range"])
        cells, found = [], {}
        for r in range(r1, (r2 or r1) + 1):
            for c in range(c1, (c2 or c1) + 1):
                v = wb[e["sheet"]].cell(row=r, column=c).value
                ref = f"{e['sheet']}!{col_let(c)}{r}"
                cells.append({"cell": ref, "value": v})
                if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) >= MIN_ABS:
                    found[ref] = hits(float(v))
        cover = defaultdict(dict)                       # sig -> {cell: (offset, operands)}
        meta = {}
        for ref, hs in found.items():
            for s_, spec, sc, sg, off, ops_ in hs:
                if ref not in cover[s_] or abs(off) < abs(cover[s_][ref][0]):
                    cover[s_][ref] = (off, ops_)
                meta[s_] = (spec, sc, sg)
        rank = lambda s_: (-len(cover[s_]), OP_RANK[meta[s_][0]["operation"]], meta[s_][1] != 1, meta[s_][2] != 1,
                           norm_label(meta[s_][0]["source"]["label"]) != norm_label(e["label"]))
        ranked = sorted(cover, key=rank)
        tested = len(found)
        best = ranked[0] if ranked else None
        rivals = [s_ for s_ in ranked if best and rank(s_) == rank(best)]
        for cell in cells:
            if cell["cell"] not in found:
                cell["result"] = "not_testable"
            elif best and cell["cell"] in cover[best]:
                off, ops_ = cover[best][cell["cell"]]
                cell.update(result="reproduced" if len(rivals) == 1 else "ambiguous", offset_months=off,
                            period=ym_str(pe + off),
                            operands=[f"{o['file']} / {o['sheet']}!{o['cell']} ({o['label']}, {o.get('role')}, {o['period']})"
                                      for o in ops_])
            else:
                cell["result"] = "not_reproduced"
        n_best = len(cover[best]) if best else 0
        status = ("not_testable" if not tested else "no_rule" if not ranked else "ambiguous" if len(rivals) > 1
                  else "reproduced_all" if n_best == tested else "reproduced_some")
        tally[status] += 1
        spec = None
        if best and len(rivals) == 1:
            sp, sc, sg = meta[best]
            spec = dict(sp, scale=sc, sign=sg)
        blocks.append({"id": e["id"], "type": e.get("type"), "status": status, "cells_tested": tested,
                       "cells_reproduced": n_best if spec else 0, "execution_spec": spec,
                       "ambiguous_between": [json.loads(s_) for s_ in rivals] if len(rivals) > 1 else [],
                       "cells": cells})
    layout = [{"layout": f.get("layout"), "roles": f.get("roles", []), "sheets": f.get("sheets", [])}
              for f in idx.get("files", []) if isinstance(f, dict)]
    json.dump({"version": 2, "workbook": os.path.basename(a.workbook), "period_end": ym_str(pe),
               "statement_layout": layout, "layout_checks": idx.get("layout_checks", []),
               "period_end_source": pe_src, "fiscal_year_start_month": fs, "blocks": blocks},
              open(a.out, "w", encoding="utf-8"), indent=1, default=str)
    print(f"MATCH SUMMARY: {len(blocks)} in-scope blocks | period end {ym_str(pe)} ({pe_src}) | fiscal year starts month {fs} | "
          + ", ".join(f"{k}={v}" for k, v in sorted(tally.items())))
    for b in blocks:
        sp = b["execution_spec"]
        rule = (f"{sp['operation']}{'(' + str(sp['months']) + ')' if sp.get('months') else ''} of "
                f"'{sp['source']['label']}' [{sp['source']['role']}] on {sp['source']['sheet']} x{sp['scale']}"
                + (" sign-flip" if sp["sign"] == -1 else "")) if sp else "-"
        offs = sorted({c.get("offset_months") for c in b["cells"] if c.get("offset_months") is not None})
        print(f"{b['id']} | {b['status']} | {b['cells_reproduced']}/{b['cells_tested']} | {rule}"
              + (f" | offsets {offs}" if offs else ""))
    for c_ in idx.get("layout_checks", []):
        print(f"LAYOUT CHECK: {c_}")
    print(f"FILE WRITTEN: {a.out}")
    print("NEXT REQUIRED STEP: write concepts.json (references/profile-fields.md), then run build_profile.py.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook"); ap.add_argument("input_map"); ap.add_argument("sources_index")
    ap.add_argument("--out", default="/app/created/matches.json"); ap.add_argument("--period-end")
    ap.add_argument("--fiscal-year-start-month", type=int)
    a = ap.parse_args()
    try:
        main(a)
    except Exception as e:
        print(f"MATCH_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
