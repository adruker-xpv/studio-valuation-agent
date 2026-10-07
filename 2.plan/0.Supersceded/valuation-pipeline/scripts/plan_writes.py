# plan_writes.py (v3) - runs each concept's EXECUTABLE rule to decide every value (deterministic)
# Every target cell has a period offset (months from the period end; 0 = current period). For each cell:
#   1. run the concept's rule for that period (value_at, sum_months, ytd_difference, components)
#   2. if the sources don't cover that period and it is an EARLIER period:
#        replay   -> carried from the workbook, reported as "historical" (not proof)
#        backtest / refresh -> shifted from the base workbook's cell for that period (the roll-forward)
#   3. otherwise the concept's missing-source rule applies (block_and_flag = exception)
# Modes: replay (same period as the profile), backtest (--actual = the analyst's real workbook), refresh.
# Usage:
#   replay:   python plan_writes.py --mode replay --workbook W --input-map M --profile P --sources-index S --out plan.json
#   backtest: python plan_writes.py --mode backtest --workbook PRIOR --actual ACTUAL --input-map M --profile P --sources-index S --period-end YYYY-MM-DD --out plan.json
#   refresh:  python plan_writes.py --mode refresh --workbook LAST --input-map M --profile P --sources-index S --period-end YYYY-MM-DD [--manual J] --out plan.json
import sys, json, argparse, os, hashlib
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from wbutil import (cells_of, ref, fingerprint, parse_period, parse_period_end, ym_to_i, ym_str, build_series, evaluate)

AUTOMATED = {"replace_period_value", "roll_window"}


def profile_hash(profile):
    return hashlib.sha256("|".join(sorted(c["approval_basis_hash"] for c in profile["concepts"])).encode()).hexdigest()[:16]


def main(a):
    imap = json.load(open(a.input_map, encoding="utf-8"))
    prof = json.load(open(a.profile, encoding="utf-8"))
    idx = json.load(open(a.sources_index, encoding="utf-8"))
    wb = load_workbook(a.workbook, data_only=False)
    actual = load_workbook(a.actual, data_only=False) if a.actual else None
    errs, notes = [], []
    fp = fingerprint(a.workbook)
    if prof.get("template_fingerprint") and prof["template_fingerprint"] != fp:
        msg = f"workbook fingerprint {fp} differs from the profile's {prof['template_fingerprint']}"
        if a.mode == "backtest":
            notes.append(msg + " (prior-period template; cell addresses assumed unchanged)")
        else:
            errs.append(msg + " (the workbook structure changed since onboarding: re-onboard with --bundle --fresh)")
    if prof.get("version") != 2:
        errs.append("profile was built by an older version; re-run build_profile.py")
    elif prof.get("status") in ("INVALID", "DRAFT"):
        errs.append(f"profile status is {prof.get('status')}; rebuild the profile")
    elif prof.get("status") == "REVIEW_REQUIRED":
        notes.append("profile has unresolved meaning (REVIEW_REQUIRED): the run tests the mechanics; resolve the open items before relying on the draft")
    cert_basis = None                     # certification gate removed: approval is recorded in company_profile.json
    if a.mode == "backtest" and not a.actual:
        errs.append("backtest needs --actual")
    if a.mode in ("backtest", "refresh") and not a.period_end:
        errs.append(f"{a.mode} needs --period-end")
    if errs:
        print("PLAN STATUS: REFUSED"); print("\n".join(f"- {e}" for e in errs))
        print("NEXT REQUIRED STEP: report these reasons to the user and stop."); return 1

    notes += [f"statement pack: {x}" for x in idx.get("layout_checks", [])]
    was = sorted({(l.get("layout"), tuple(l.get("roles", []))) for l in prof.get("statement_layout", [])})
    now = sorted({(f.get("layout"), tuple(f.get("roles", []))) for f in idx.get("files", []) if isinstance(f, dict)})
    if was and now and was != now:
        notes.append(f"statement layout differs from the profile's: was {was}, now {now}")
    fper = {f["name"]: ym_to_i(parse_period(f["statement_period"])) for f in idx.get("files", [])
            if isinstance(f, dict) and f.get("statement_period")}
    series = build_series(idx["values"], fper)
    pe = ym_to_i(parse_period_end(a.period_end)) if a.period_end else ym_to_i(parse_period(prof["period_end"]))
    manual = json.load(open(a.manual, encoding="utf-8")) if a.manual else {}
    manual = manual.get("values", manual)
    plan = []

    def val(book, k):
        return book[k[0]].cell(row=k[1], column=k[2]).value if book else None

    def add(k, c, value, status, origin, msg="", off=None):
        plan.append({"cell": ref(k), "concept_id": c["concept_id"], "write_mode": c["write_mode"],
                     "period": ym_str(pe + off) if off is not None else None, "offset_months": off,
                     "value": value, "base_value": val(wb, k), "actual_value": val(actual, k) if actual else None,
                     "status": status, "origin": origin, "message": msg})

    def missing(k, c, msg, off):
        mb = c.get("missing_source_behavior", "block_and_flag")
        if mb == "carry_forward_and_flag":
            add(k, c, val(wb, k), "flag", {"type": "carried_forward_missing_source"}, msg + "; carried forward", off)
        elif mb == "manual_input" and ref(k) in manual:
            add(k, c, manual[ref(k)], "flag", {"type": "manual", "provided_by": "manual_inputs.json"}, msg, off)
        else:
            add(k, c, None, "exception", {"type": "missing"}, f"{msg} ({mb})", off)

    for c in prof["concepts"]:
        cells = [k for t in c["targets"] for k in cells_of(t)]
        wm = c["write_mode"]
        if wm in ("carry_forward", "carry_forward_event_driven"):
            for k in cells:
                add(k, c, val(wb, k), "flag" if wm == "carry_forward_event_driven" and a.mode != "replay" else "ok",
                    {"type": "carry_forward"},
                    "carried forward; confirm no capital or instrument event" if wm == "carry_forward_event_driven" else "")
            continue
        if wm == "manual_input":
            for k in cells:
                if a.mode == "replay":
                    add(k, c, val(wb, k), "ok", {"type": "manual", "provided_by": "base workbook (replay)"})
                elif a.mode == "backtest":
                    add(k, c, val(actual, k), "ok", {"type": "manual", "provided_by": "actual workbook (backtest)"})
                elif ref(k) in manual:
                    add(k, c, manual[ref(k)], "ok", {"type": "manual", "provided_by": "manual_inputs.json"})
                elif c.get("missing_source_behavior") == "carry_forward_and_flag":
                    add(k, c, val(wb, k), "flag", {"type": "carry_forward"},
                        "no new value given: last period's value carried forward - confirm")
                else:
                    add(k, c, None, "exception", {"type": "missing"}, "manual input not provided")
            continue
        if wm not in AUTOMATED or not c.get("execution"):
            for k in cells:
                add(k, c, None, "exception", {"type": "no_rule"}, f"{c['status']}: no executable rule for write_mode {wm}")
            continue
        offs = c.get("cell_offsets_months", {})
        by_off = {v: r_ for r_, v in offs.items() if v is not None}
        for k in cells:
            o = offs.get(ref(k))
            if o is None:
                add(k, c, val(wb, k), "historical", {"type": "historical_constant"},
                    "no period mapped to this cell; kept from the workbook")
                continue
            v, ops_, why = evaluate(c["execution"], series, pe + o)
            if v is not None:
                add(k, c, v, "ok", {"type": "rule", "operation": c["execution"]["operation"],
                                    "operands": [f"{x['file']} / {x['sheet']}!{x['cell']} ({x.get('role')}, {x['period']})"
                                                 for x in ops_]}, "", o)
                continue
            if o < 0 and a.mode == "replay":
                add(k, c, val(wb, k), "historical", {"type": "historical_not_in_sources"},
                    f"{why}; earlier period kept from the workbook (not rebuilt)", o)
            elif o < 0:
                src_cell = by_off.get(o + a.advance_months)
                if src_cell:
                    sh, cc = src_cell.rsplit("!", 1)
                    add(k, c, wb[sh][cc].value, "ok", {"type": "shifted", "from": src_cell},
                        f"rolled forward from {src_cell} in the base workbook", o)
                else:
                    missing(k, c, f"{why}; no base cell to roll forward from", o)
            else:
                missing(k, c, why, o)

    for p in plan:
        sh, rest = p["cell"].rsplit("!", 1)
        if wb[sh][rest].data_type == "f":
            p.update(status="exception", value=None, message="target is a formula cell; refusing to write")
    count = lambda s: sum(1 for p in plan if p["status"] == s)
    ex = [p for p in plan if p["status"] == "exception"]
    carried_manual = [p for p in plan if p["write_mode"] == "manual_input" and p["status"] == "flag"]
    status = "BLOCKED" if ex else ("INPUT_REQUIRED" if carried_manual else "READY")
    json.dump({"version": 3, "mode": a.mode, "workbook": os.path.basename(a.workbook),
               "actual": os.path.basename(a.actual) if a.actual else None, "period_end": ym_str(pe),
               "advance_months": a.advance_months, "template_fingerprint": fp, "profile_hash": profile_hash(prof),
               "certification_basis": cert_basis, "profile_status": prof.get("status"), "notes": notes, "status": status,
               "writes": plan},
              open(a.out, "w", encoding="utf-8"), indent=1, default=str)
    print(f"PLAN STATUS: {status} | mode: {a.mode} | period end: {ym_str(pe)} | cells: {len(plan)} | rebuilt by rule: "
          f"{sum(1 for p in plan if p['origin'].get('type') == 'rule')} | rolled forward: "
          f"{sum(1 for p in plan if p['origin'].get('type') == 'shifted')} | historical kept: {count('historical')} | "
          f"flagged: {count('flag')} | exceptions: {len(ex)}")
    for n in notes:
        print(f"note: {n}")
    for p in [p for p in plan if p["status"] in ("exception", "flag", "historical")][:60]:
        print(f"- {p['status'].upper()} {p['cell']} [{p['concept_id']}]: {p['message']}")
    print(f"FILE WRITTEN: {a.out}")
    print("NEXT REQUIRED STEP: " + ("stop: report the exceptions (they mean a rule could not run) and do NOT continue to "
                                    "Excel." if ex else "run apply_writes.py."))
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--mode", choices=["replay", "backtest", "refresh"], required=True)
    ap.add_argument("--workbook", required=True); ap.add_argument("--input-map", required=True)
    ap.add_argument("--profile", required=True); ap.add_argument("--sources-index", required=True)
    ap.add_argument("--manual"); ap.add_argument("--certification"); ap.add_argument("--actual")
    ap.add_argument("--period-end"); ap.add_argument("--advance-months", type=int, default=3)
    ap.add_argument("--matches", help="ignored (kept for compatibility)")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    try:
        sys.exit(main(a))
    except Exception as e:
        print(f"PLAN_ERROR: {type(e).__name__}: {e}"); sys.exit(1)
