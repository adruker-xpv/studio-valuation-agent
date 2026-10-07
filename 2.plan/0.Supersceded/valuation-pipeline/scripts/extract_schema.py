# extract_schema.py - deterministic Excel workbook schema extractor (v4)
# Usage:
#   python extract_schema.py <workbook.xlsx> [--out DIR] [--key-outputs-file key_outputs.json]
#   (legacy: --key-outputs "Sheet!B3,'My Sheet'!C10:C12")
# Prints a bounded OVERVIEW. Writes two complete, never-truncated files to --out:
#   <stem>_schema.md              overview + sheet index with line numbers + outputs grouped by
#                                 sheet + every sheet section (one file, nothing cut)
#   <stem>_inputs_skeleton.json   every input block (with dependency paths), for the agent to annotate
# Only needs openpyxl. Never modifies the workbook.
MAX_KEY_OUTPUTS = 25                 # tracing is cheap Python; the cap only bounds review
import os, re, glob, json, argparse, datetime, traceback, hashlib
from collections import defaultdict, deque
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx, get_column_letter as col_let
from openpyxl.utils import range_boundaries

MAX_RANGE_CELLS = 20000     # larger ranges are resolved against existing cells only
OVERVIEW_LIST_CAP = 60      # max list lines in the printed overview (files are never capped)

STR_LIT = re.compile(r'"[^"]*"')
REF = re.compile(
    r"(?<![A-Za-z0-9_.])"
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!)?"
    r"(\$?)([A-Z]{1,3})(\$?)([0-9]+)"
    r"(?::(\$?)([A-Z]{1,3})(\$?)([0-9]+))?"
    r"(?![A-Za-z0-9_(])")
SHEET_PFX = r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_.]*))!)?"
COLREF = re.compile(r"(?<![A-Za-z0-9_.$])" + SHEET_PFX + r"\$?([A-Z]{1,3}):\$?([A-Z]{1,3})(?![A-Za-z0-9_(])")
ROWREF = re.compile(r"(?<![A-Za-z0-9_.$:])" + SHEET_PFX + r"\$?([0-9]+):\$?([0-9]+)(?![A-Za-z0-9_(:])")
MAX_ROW, MAX_COL = 1048576, 16384
NUM = re.compile(r"(?<![A-Za-z0-9_.\[\]$])(\d+(?:\.\d+)?%?)(?![A-Za-z0-9_\[\]])")
KEY_REF = re.compile(r"(?:'((?:[^']|'')+)'|([^!,']+))!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)")
MONEY_WORDS = re.compile(r"amount|value|price|revenue|cost|cash|debt|balance|ebitda|proceeds|\$|total", re.I)
DATE_WORDS = re.compile(r"\bdate\b|as of|as at|maturity date|issue date", re.I)


def find_file():
    for root in ["/app", os.getcwd(), "/mnt/data", "/tmp"]:
        hits = [h for ext in ("xlsx", "xlsm")
                for h in glob.glob(os.path.join(root, "**", "*." + ext), recursive=True)]
        hits = [h for h in hits if not os.path.basename(h).startswith("~$") and "/created/" not in h]
        if hits:
            return sorted(hits)[0]
    raise FileNotFoundError("no .xlsx or .xlsm file found in sandbox")


def parse_refs(formula, cur_sheet):
    out = []
    for m in REF.finditer(STR_LIT.sub('""', formula)):
        sheet = m.group(1) or m.group(2)
        sheet = sheet.replace("''", "'") if sheet else cur_sheet
        c1, r1 = col_idx(m.group(4)), int(m.group(6))
        c2, r2 = (col_idx(m.group(8)), int(m.group(10))) if m.group(8) else (c1, r1)
        out.append((sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)))
    clean = STR_LIT.sub('""', formula)
    for m in COLREF.finditer(clean):          # whole columns, e.g. A:A, Sheet!B:D
        sheet = (m.group(1) or m.group(2) or cur_sheet).replace("''", "'")
        c1, c2 = sorted((col_idx(m.group(3)), col_idx(m.group(4))))
        out.append((sheet, 1, c1, MAX_ROW, c2))
    for m in ROWREF.finditer(clean):          # whole rows, e.g. 3:3, Sheet!5:9
        sheet = (m.group(1) or m.group(2) or cur_sheet).replace("''", "'")
        r1, r2 = sorted((int(m.group(3)), int(m.group(4))))
        out.append((sheet, r1, 1, r2, MAX_COL))
    return out


def rel_key(formula, row, col):
    def one(ca, c, ra, r):
        ci, ri = col_idx(c), int(r)
        return (f"R{ri}" if ra else f"R[{ri-row}]") + (f"C{ci}" if ca else f"C[{ci-col}]")

    def sub(m):
        s = m.group(1) or m.group(2) or ""
        s = s + "!" if s else ""
        s += one(m.group(3), m.group(4), m.group(5), m.group(6))
        if m.group(8):
            s += ":" + one(m.group(7), m.group(8), m.group(9), m.group(10))
        return s
    return REF.sub(sub, STR_LIT.sub('""', formula))


def hardcodes(formula):
    f = REF.sub("", STR_LIT.sub('""', formula))
    return sorted({n for n in NUM.findall(f) if n not in ("0", "1")})


def fmt(v):
    if v is None:
        return "(blank)"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == datetime.time(0) else v.isoformat(sep=" ")
    if isinstance(v, (int, float)):
        return f"{v:,.0f}" if abs(v) >= 1000 else f"{v:.4g}"
    s = str(v).replace("\n", " ").strip()
    return s[:40] + ("..." if len(s) > 40 else "")


def addr(r, c1, c2=None):
    a = f"{col_let(c1)}{r}"
    return a + (f":{col_let(c2)}{r}" if c2 and c2 != c1 else "")


def qs(sheet):
    return f"'{sheet}'" if re.search(r"[^A-Za-z0-9_]", sheet) else sheet


def main(path, out_dir, key_outputs_arg, ko_reasons=None, ko_controls=None):
    ko_reasons = ko_reasons or {}
    ko_controls = ko_controls or {}
    wbf = load_workbook(path, data_only=False)
    wbv = load_workbook(path, data_only=True)
    sheetnames = wbf.sheetnames

    # ---------- collect cells ----------
    formulas, consts = {}, {}
    texts = defaultdict(dict)          # label lookup text (merged-cell aware)
    sheet_cells = defaultdict(list)
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                key = (ws.title, cell.row, cell.column)
                sheet_cells[ws.title].append((cell.row, cell.column))
                ftxt = str(getattr(v, "text", v))
                if cell.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                    formulas[key] = ftxt if ftxt.startswith("=") else "=" + ftxt
                else:
                    consts[key] = v
                    if isinstance(v, str) and v.strip():
                        texts[ws.title][(cell.row, cell.column)] = v.strip()
        for mr in ws.merged_cells.ranges:   # spread merged header text across its range
            t = texts[ws.title].get((mr.min_row, mr.min_col))
            if t:
                for rr in range(mr.min_row, mr.max_row + 1):
                    for cc in range(mr.min_col, mr.max_col + 1):
                        texts[ws.title].setdefault((rr, cc), t)

    # ---------- named ranges ----------
    items = []
    try:
        items = list(wbf.defined_names.items())
    except Exception:
        items = [(d.name, d) for d in getattr(wbf.defined_names, "definedName", [])]
    for ws in wbf.worksheets:
        try:
            items += [(n, d) for n, d in ws.defined_names.items()]
        except Exception:
            pass
    names = {}
    for n, d in items:
        try:
            dests = list(d.destinations)
        except Exception:
            dests = []
        names[n] = {"refers_to": getattr(d, "attr_text", "") or "", "dests": dests}
    name_re = None
    if names:
        name_re = re.compile(r"(?<![A-Za-z0-9_.])(" + "|".join(
            sorted(map(re.escape, names), key=len, reverse=True)) + r")(?![A-Za-z0-9_(])", re.I)
    lower_names = {n.lower(): n for n in names}
    used_names = set()

    # ---------- dependency graph ----------
    prec = defaultdict(set)            # formula cell -> cells it reads
    referenced = set()
    referenced_single = set()          # cells referenced one-by-one (not via a multi-cell range)
    edges = defaultdict(int)
    external = set()

    def expand(rs, r1, c1, r2, c2):
        if (r2 - r1 + 1) * (c2 - c1 + 1) <= MAX_RANGE_CELLS:
            return [(rs, rr, cc) for rr in range(r1, r2 + 1) for cc in range(c1, c2 + 1)]
        return [(rs, rr, cc) for (rr, cc) in sheet_cells.get(rs, []) if r1 <= rr <= r2 and c1 <= cc <= c2]

    for key, f in formulas.items():
        s = key[0]
        if re.search(r"\[\d+\]", f) or ".xls" in f.lower():
            external.add(key)
        refs = list(parse_refs(f, s))
        if name_re:
            for m in name_re.finditer(STR_LIT.sub('""', f)):
                n = lower_names[m.group(1).lower()]
                used_names.add(n)
                for dsheet, coord in names[n]["dests"]:
                    try:
                        c1, r1, c2, r2 = range_boundaries(coord.replace("$", ""))
                        c1, r1 = c1 or 1, r1 or 1
                        c2 = c2 or (MAX_COL if coord.replace("$", "")[0].isdigit() else c1)
                        r2 = r2 or (MAX_ROW if not coord.replace("$", "")[0].isdigit() else r1)
                        refs.append((dsheet, r1, c1, r2, c2))
                    except Exception:
                        pass
        for (rs, r1, c1, r2, c2) in refs:
            if rs != s:
                edges[(rs, s)] += 1
            cells = expand(rs, r1, c1, r2, c2)
            if r1 == r2 and c1 == c2:
                referenced_single.add((rs, r1, c1))
            prec[key].update(cells)
            referenced.update(cells)

    for n, info in names.items():
        t = info["refers_to"]
        if n.lower().startswith("_xlnm.") or n.lower() in ("print_area", "print_titles"):
            cat = "print_or_builtin"
        elif "#REF" in t.upper():
            cat = "broken"
        elif re.search(r"\[\d+\]", t) or ".xls" in t.lower():
            cat = "external"
        elif n in used_names:
            cat = "referenced"
        else:
            cat = "unused"
        info["category"] = cat

    def val(k):
        return wbv[k[0]].cell(row=k[1], column=k[2]).value

    # ---------- outputs and materiality ----------
    output_cells = sorted(k for k in formulas if k not in referenced)
    key_outputs, key_errors = [], []
    if key_outputs_arg:
        for m in KEY_REF.finditer(key_outputs_arg):
            sh = (m.group(1) or m.group(2)).strip().replace("''", "'")
            if sh not in sheetnames:
                key_errors.append(f"unknown sheet in key outputs: {sh}")
                continue
            c1, r1, c2, r2 = range_boundaries(m.group(3).replace("$", ""))
            c2, r2 = c2 or c1, r2 or r1
            for rr in range(r1, (r2 or r1) + 1):
                for cc in range(c1, (c2 or c1) + 1):
                    if (sh, rr, cc) in formulas or (sh, rr, cc) in consts:
                        key_outputs.append((sh, rr, cc))
                    else:
                        key_errors.append(f"key output is an empty cell: {sh}!{addr(rr, cc)}")
    if key_errors:
        raise ValueError("; ".join(key_errors))

    def spec_for(k, table):
        for raw, spec in table.items():
            for m in KEY_REF.finditer(raw):
                sh = (m.group(1) or m.group(2)).strip()
                c1, r1, c2, r2 = range_boundaries(m.group(3).replace("$", ""))
                if sh == k[0] and r1 <= k[1] <= (r2 or r1) and c1 <= k[2] <= (c2 or c1):
                    return spec
        return None

    def reason_for(k):
        for raw, why in ko_reasons.items():
            for m in KEY_REF.finditer(raw if "'" in raw else raw):
                sh = (m.group(1) or m.group(2)).strip()
                c1, r1, c2, r2 = range_boundaries(m.group(3).replace("$", ""))
                if sh == k[0] and r1 <= k[1] <= (r2 or r1) and c1 <= k[2] <= (c2 or c1):
                    return why
        return ""

    def backward(starts):
        seen, q = set(starts), deque(starts)
        while q:
            k = q.popleft()
            for p in prec.get(k, ()):
                if p not in seen:
                    seen.add(p)
                    q.append(p)
        return seen

    targets = key_outputs or output_cells
    reach = backward(targets)
    def ref(k):
        return f"{k[0]}!{addr(k[1], k[2])}"

    feeds = defaultdict(set)
    parents = {}                       # key output name -> {cell: next cell toward that output}
    if key_outputs:
        for ko in key_outputs[:50]:
            par, q = {ko: None}, deque([ko])
            while q:
                k = q.popleft()
                for p in prec.get(k, ()):
                    if p not in par:
                        par[p] = k
                        q.append(p)
            kn = ref(ko)
            parents[kn] = par
            for k in par:
                feeds[k].add(kn)

    def path_to(cell, kn):
        par, out = parents[kn], []
        while cell is not None:
            out.append(ref(cell))
            cell = par[cell]
        return out

    # ---------- labels / grouping ----------
    def label(s, r, c):
        t = texts.get(s, {})
        left = next((t[(r, cc)] for cc in range(c - 1, 0, -1) if (r, cc) in t), "")
        up = next((t[(rr, c)] for rr in range(r - 1, max(0, r - 40), -1) if (rr, c) in t), "")
        return (fmt(left) if left else "?"), (fmt(up) if up else "")

    def group_row(cells, keyfn):
        groups = []
        for (r, c) in sorted(cells):
            k = keyfn(r, c)
            g = groups[-1] if groups else None
            if g and g["r"] == r and g["c2"] == c - 1 and g["k"] == k:
                g["c2"] = c
                g["cells"].append((r, c))
            else:
                groups.append({"r": r, "c1": c, "c2": c, "k": k, "cells": [(r, c)]})
        return groups

    def type_flags(lbl, hdr, values):
        fl = []
        text = f"{lbl} {hdr}"
        if MONEY_WORDS.search(text) and any(isinstance(v, (datetime.date, datetime.datetime)) for v in values):
            fl.append("label_type_mismatch: money-like label but date value")
        if DATE_WORDS.search(lbl) and any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values):
            fl.append("label_type_mismatch: date-like label but numeric value")
        return fl

    skeleton, sheet_texts, sheet_summ = [], {}, []
    flags_all = defaultdict(list)       # category -> list of strings
    blocking = []
    for ws in wbf.worksheets:
        s = ws.title
        f_cells = [(r, c) for (ss, r, c) in formulas if ss == s]
        # text reached only through multi-cell ranges (e.g. a header inside SUM(A:A)) is a label, not an input
        in_cells = [(r, c) for (ss, r, c), v in consts.items() if ss == s and (s, r, c) in referenced
                    and (not isinstance(v, str) or (s, r, c) in referenced_single)]
        unref_nums = sum(1 for (ss, r, c), v in consts.items()
                         if ss == s and isinstance(v, (int, float)) and (s, r, c) not in referenced)
        lines = []
        # inputs
        in_lines, mat_cells = [], 0
        for g in group_row(in_cells, lambda r, c: ((s, r, c) in reach)):
            keys = [(s, r, c) for (r, c) in g["cells"]]
            material = (s, g["r"], g["c1"]) in reach
            mat_cells += len(keys) if material else 0
            left, up1 = label(s, g["r"], g["c1"])
            up2 = label(s, g["r"], g["c2"])[1]
            hdr = up1 if up1 == up2 or not up2 else f"{up1}..{up2}"
            raw = [consts[k] for k in keys]
            vals = [fmt(v) for v in raw]
            rng = addr(g["r"], g["c1"], g["c2"])
            fl = type_flags(left, hdr, raw)
            fed = sorted(set().union(*[feeds.get(k, set()) for k in keys])) if key_outputs else []
            paths = {}
            for kn in fed:
                cands = [path_to(k, kn) for k in keys if k in parents[kn]]
                paths[kn] = min(cands, key=len)
            vtxt = ", ".join(vals[:6]) + (f", ... ({len(vals)} cells)" if len(vals) > 6 else "")
            line = (f"{rng} [input] {left}" + (f" | {hdr}" if hdr else "") + f" = {vtxt}"
                    + ("  MATERIAL" if material else "  not-material"))
            if fed:
                line += f"  feeds: {', '.join(fed[:5])}" + (f" +{len(fed) - 5}" if len(fed) > 5 else "")
                shortest = min(paths.values(), key=len)
                line += f"  path: {' -> '.join(shortest)}"
            for x in fl:
                line += f"  FLAG:{x}"
                flags_all["label_type_mismatch"].append(f"{s}!{rng}")
            in_lines.append(line)
            skeleton.append({"id": f"{s}!{rng}", "sheet": s, "range": rng, "label": left, "header": hdr,
                             "n_cells": len(vals), "sample_values": vals[:6], "material": material,
                             "feeds_key_outputs": fed, "paths_to_key_outputs": paths, "flags": fl,
                             "type": None, "expected_source": None, "changes_each_quarter": None,
                             "notes": None})
        # formulas
        f_lines, row_groups, n_out = [], defaultdict(list), 0
        for g in group_row(f_cells, lambda r, c: rel_key(formulas[(s, r, c)], r, c)):
            row_groups[g["r"]].append(g)
            r, c1 = g["r"], g["c1"]
            keys = [(s, rr, cc) for (rr, cc) in g["cells"]]
            is_out = not any(k in referenced for k in keys)
            is_key = any(k in key_outputs for k in keys)
            n_out += len(keys) if is_out else 0
            left, up1 = label(s, r, c1)
            up2 = label(s, r, g["c2"])[1]
            hdr = up1 if up1 == up2 or not up2 else f"{up1}..{up2}"
            vals = [val(k) for k in keys]
            vtxt = ", ".join(fmt(v) for v in vals[:3]) + (", ..." if len(vals) > 3 else "")
            n = len(keys)
            role = "KEY OUTPUT" if is_key else ("output" if is_out else "calc")
            rng = addr(r, c1, g["c2"])
            line = (f"{rng} [{role}] {left}" + (f" | {hdr}" if hdr else "")
                    + f" {formulas[keys[0]]}" + (f"  (same pattern x{n})" if n > 1 else "")
                    + f" -> {vtxt}")
            hc = hardcodes(formulas[keys[0]])
            if hc:
                line += f"  FLAG:hardcoded {','.join(hc)}"
                flags_all["hardcoded_in_formula"].append(f"{s}!{rng} ({','.join(hc)})")
            errs = [v for v in vals if isinstance(v, str) and v.startswith("#")]
            if errs:
                line += f"  FLAG:error {errs[0]}"
                flags_all["error"].append(f"{s}!{rng} {errs[0]}")
                if any(k in reach for k in keys):
                    blocking.append(f"error {errs[0]} in {s}!{rng} feeds "
                                    + ("key outputs" if key_outputs else "workbook outputs"))
            if any(k in external for k in keys):
                line += "  FLAG:external_link"
                flags_all["external_link"].append(f"{s}!{rng}")
                if any(k in reach for k in keys):
                    blocking.append(f"external workbook link in {s}!{rng} feeds "
                                    + ("key outputs" if key_outputs else "workbook outputs"))
            if all(v is None for v in vals):
                line += "  FLAG:no cached value"
                flags_all["no_cached_value"].append(f"{s}!{rng}")
            f_lines.append(line)
        for r, gs in row_groups.items():
            for i in range(1, len(gs) - 1):
                a, b, c_ = gs[i - 1], gs[i], gs[i + 1]
                if (a["k"] == c_["k"] and len(b["cells"]) == 1
                        and a["c2"] + 1 == b["c1"] and b["c2"] + 1 == c_["c1"]):
                    flags_all["pattern_break"].append(f"{s}!{addr(r, b['c1'])}")

        reads = sorted({a for (a, b) in edges if b == s})
        read_by = sorted({b for (a, b) in edges if a == s})
        header = (f"state: {ws.sheet_state} | used range: {ws.dimensions} | formulas: {len(f_cells)}"
                  f" | input cells: {len(in_cells)} (material: {mat_cells}) | output cells: {n_out}"
                  f" | unreferenced numbers: {unref_nums}")
        lines = [f"=== SHEET: {s} ===", header,
                 f"reads from: {', '.join(reads) or '-'} | read by: {', '.join(read_by) or '-'}",
                 "-- INPUTS --"] + (in_lines or ["(none)"]) + ["-- FORMULAS --"] + (f_lines or ["(none)"]) + \
                ["=== END OF SHEET ==="]
        sheet_texts[s] = "\n".join(lines)
        sheet_summ.append((s, ws.sheet_state, len(f_cells), len(in_cells), mat_cells, n_out))

    # ---------- write files (two files only) ----------
    os.makedirs(out_dir, exist_ok=True)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", os.path.splitext(os.path.basename(path))[0])

    # outputs grouped by sheet, copied-across rows collapsed into one line
    by_sheet = defaultdict(list)
    for k in output_cells:
        by_sheet[k[0]].append((k[1], k[2]))
    out_sec = ["=== OUTPUTS BY SHEET (formulas nothing else uses; KEY OUTPUT candidates) ==="]
    for sname in sheetnames:
        cells = by_sheet.get(sname)
        if not cells:
            continue
        out_sec.append(f"-- {sname} ({len(cells)} cells) --")
        for g in group_row(cells, lambda r, c: rel_key(formulas[(sname, r, c)], r, c)):
            vals = [fmt(val((sname, rr, cc))) for (rr, cc) in g["cells"]]
            left, up1 = label(sname, g["r"], g["c1"])
            up2 = label(sname, g["r"], g["c2"])[1]
            hdr = up1 if up1 == up2 or not up2 else f"{up1}..{up2}"
            up = backward([(sname, g["r"], g["c1"])])
            n_up = sum(1 for k in up if k in consts)
            n_sh = len({k[0] for k in up if k in consts})
            out_sec.append(f"{qs(sname)}!{addr(g['r'], g['c1'], g['c2'])} {left}" + (f" | {hdr}" if hdr else "")
                           + f" = {', '.join(vals[:4])}" + (f", ... ({len(vals)} cells)" if len(vals) > 4 else "")
                           + f"  [upstream inputs: {n_up} from {n_sh} sheets]")
    out_sec.append("=== END OF OUTPUTS ===")

    n_blocks = len(skeleton)
    n_cells = sum(e["n_cells"] for e in skeleton)
    m_blocks = sum(1 for e in skeleton if e["material"])
    m_cells = sum(e["n_cells"] for e in skeleton if e["material"])
    basis = "confirmed key outputs" if key_outputs else "ALL output cells (key outputs NOT set)"

    def capped(xs):
        return xs[:OVERVIEW_LIST_CAP] + ([f"... {len(xs) - OVERVIEW_LIST_CAP} more in the schema file"]
                                         if len(xs) > OVERVIEW_LIST_CAP else [])

    def build_overview(index_lines):
        ov = ["=== WORKBOOK OVERVIEW ===",
              f"file: {os.path.basename(path)} | sheets: {len(sheetnames)} | formula cells: {len(formulas)}",
              f"input blocks: {n_blocks} | input cells: {n_cells} | material blocks: {m_blocks}"
              f" | material cells: {m_cells} | output cells: {len(output_cells)}",
              f"materiality basis: {basis}",
              "legend: [input]=hardcoded value a formula uses; [calc]=formula other formulas use; "
              "[output]=formula nothing uses; MATERIAL=input that feeds the materiality basis; "
              "path=shortest dependency chain from the input to a key output; "
              "values after -> are Excel's last saved results; labels are guesses from nearby text",
              "completeness: the schema file holds every section in full; overview lists may be capped",
              "", "=== INDEX (line ranges in the schema file; read each range to its END line) ==="]
        ov += index_lines
        ov += ["", "=== SHEET FLOW (source -> uses it: reference count) ==="]
        ov += [f"{a} -> {b}: {n}" for (a, b), n in sorted(edges.items(), key=lambda x: -x[1])] or ["(none)"]
        ov += ["", "=== NAMED RANGES (category: name = refers to) ==="]
        ov += capped([f"{i['category']}: {n} = {i['refers_to']}" for n, i in sorted(
            names.items(), key=lambda x: (x[1]["category"], x[0]))]) or ["(none)"]
        ov += ["", "=== KEY OUTPUTS ==="]
        ov += ([f"{ref(k)} = {fmt(val(k))}  {label(k[0], k[1], k[2])[0]}" for k in key_outputs]
               or [f"NOT SET. Candidates: OUTPUTS BY SHEET section ({len(output_cells)} output cells)."])
        ov += ["", "=== BLOCKING ISSUES (errors or external links that feed the materiality basis) ==="]
        ov += capped(sorted(set(blocking))) or ["(none)"]
        ov += ["", "=== FLAG COUNTS (full detail in the sheet sections) ==="]
        ov += [f"{k}: {len(v)} (e.g. {', '.join(v[:5])})" for k, v in sorted(flags_all.items())] or ["(none)"]
        ov += ["", "=== END OF OVERVIEW ==="]
        return ov

    sections = [("OUTPUTS BY SHEET", out_sec, "")] + [
        (s_, sheet_texts[s_].split("\n"), next(
            f"formulas {nf}, input cells {ni} (material {nm}), outputs {no}"
            + (f", {st}" if st != "visible" else "") for (x, st, nf, ni, nm, no) in sheet_summ if x == s_))
        for s_ in sheetnames]
    placeholder = [""] * len(sections)
    ov_len = len(build_overview(placeholder))
    index, line_no = [], ov_len + 2           # one blank line after the overview
    for name, body, summ in sections:
        index.append(f"{name}: lines {line_no}-{line_no + len(body) - 1}" + (f" | {summ}" if summ else ""))
        line_no += len(body) + 1              # blank line between sections
    overview = build_overview(index)
    doc = "\n".join(overview) + "\n\n" + "\n\n".join("\n".join(b) for _, b, _ in sections) + "\n"

    schema_path = os.path.join(out_dir, f"{stem}_schema.md")
    with open(schema_path, "w", encoding="utf-8") as fh:
        fh.write(doc)
    controls = []
    for k in key_outputs:
        c = spec_for(k, ko_controls)
        if not c:
            continue
        v = val(k)
        ctype, exp = c.get("control_type"), c.get("expected")
        tol = float(c.get("absolute_tolerance", 0) or 0)
        if isinstance(v, str) and v.startswith("#"):
            result = "FAIL"
        elif ctype == "boolean":
            result = "PASS" if (v is True or v is False) and v == bool(exp) else "FAIL"
        elif ctype == "numeric":
            ok = isinstance(v, (int, float)) and not isinstance(v, bool) and isinstance(exp, (int, float))
            result = "PASS" if ok and abs(v - exp) <= tol else "FAIL"
        else:
            result = "INVALID_SPEC"
        controls.append({"cell": ref(k), "control_type": ctype, "expected": exp, "absolute_tolerance": tol,
                         "severity": c.get("severity", "blocking"), "value": fmt(v), "result": result})
    fp_src = "|".join(sheetnames) + "||" + "|".join(sorted(
        f"{k[0]}!{rel_key(f, k[1], k[2])}@{k[1]},{k[2]}" for k, f in formulas.items()))
    template_fp = hashlib.sha256(fp_src.encode("utf-8")).hexdigest()[:16]
    skel_path = os.path.join(out_dir, f"{stem}_inputs_skeleton.json")
    with open(skel_path, "w", encoding="utf-8") as fh:
        json.dump({"workbook": os.path.basename(path), "generated_by": "extract_schema.py v4",
                   "key_outputs": [ref(k) for k in key_outputs],
                   "key_output_reasons": {ref(k): reason_for(k) for k in key_outputs},
                   "controls": controls,
                   "key_output_formula_precedents": {kn: sorted(ref(k) for k in par if k in formulas)
                                                     for kn, par in parents.items()},
                   "template_fingerprint": template_fp,
                   "materiality_basis": basis,
                   "totals": {"blocks": n_blocks, "cells": n_cells,
                              "material_blocks": m_blocks, "material_cells": m_cells},
                   "named_ranges": [{"name": n, "refers_to": i["refers_to"], "category": i["category"],
                                     "used_by_formulas": n in used_names}
                                    for n, i in sorted(names.items())],
                   "blocking_issues": sorted(set(blocking)),
                   "inputs": skeleton}, fh, indent=1, default=str)
    print("\n".join(overview))
    if controls:
        print("\n=== CONTROLS (cached values from the last Excel save) ===")
        for c in controls:
            print(f"{c['cell']}: {c['result']} (value {c['value']}, expected {c['expected']}"
                  f"{', tolerance ' + str(c['absolute_tolerance']) if c['control_type'] == 'numeric' else ''}, {c['severity']})")
    print(f"template fingerprint: {template_fp}")
    print(f"\nFILES WRITTEN:\n{schema_path}\n{skel_path}")
    if not key_outputs:
        print("\nNEXT REQUIRED STEP (J1): choose key outputs yourself using references/key-output-rules.md, "
              "write /app/created/key_outputs.json, then re-run this script with --key-outputs-file. "
              "Do not ask the user. Do not report results yet.")
    else:
        print("\nNEXT REQUIRED STEP (J2): read the sheet sections with material inputs, then write "
              "/app/created/classifications.json for every MATERIAL block (see references/classification-rules.md) "
              f"and run apply_classifications.py. The skeleton is NOT an input map. Do not report results yet.")

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", help="path to .xlsx/.xlsm (searched for if omitted)")
    ap.add_argument("--out", default="/app/created" if os.path.isdir("/app") else os.getcwd())
    ap.add_argument("--key-outputs", default="", help="comma-separated Sheet!A1 refs (legacy)")
    ap.add_argument("--key-outputs-file", default="", help='JSON: {"key_outputs": [{"cell": "Sheet!B5", "reason": "..."}]}')
    a = ap.parse_args()
    try:
        ko_arg, reasons = a.key_outputs, {}
        if a.key_outputs_file:
            spec = json.load(open(a.key_outputs_file, encoding="utf-8"))
            items = spec.get("key_outputs", [])
            if items and isinstance(items[0], str):          # reuse from a saved input map
                rs = spec.get("key_output_reasons", {})
                items = [{"cell": c, "reason": rs.get(c) or "reused from saved input map"} for c in items]
            if not items or len(items) > MAX_KEY_OUTPUTS:
                raise ValueError(f"key_outputs.json must list 1-{MAX_KEY_OUTPUTS} key outputs")
            for it in items:
                if not it.get("cell") or not it.get("reason"):
                    raise ValueError(f"each key output needs 'cell' and 'reason': {it}")
            ko_arg = ",".join(it["cell"] for it in items)
            reasons = {it["cell"].replace("'", ""): it["reason"] for it in items}
            ctrls = {it["cell"].replace("'", ""): it["control"] for it in items if it.get("control")}
            for c in spec.get("controls", []):                  # reuse controls from a saved input map
                ctrls.setdefault(c["cell"], c)
            for cell, c in ctrls.items():
                if c.get("control_type") not in ("boolean", "numeric") or "expected" not in c \
                        or c.get("severity", "blocking") not in ("blocking", "warning"):
                    raise ValueError(f"invalid control spec for {cell}: needs control_type boolean|numeric, "
                                     f"expected, severity blocking|warning")
        else:
            ctrls = {}
        main(a.workbook or find_file(), a.out, ko_arg, reasons, ctrls)
    except Exception as e:
        frames = traceback.extract_tb(e.__traceback__)
        own = [f for f in frames if f.filename.endswith("extract_schema.py")]
        where = f"extract_schema.py line {own[-1].lineno} in {own[-1].name}" if own else "startup"
        if frames and not frames[-1].filename.endswith("extract_schema.py"):
            where += f", raised inside {os.path.basename(frames[-1].filename)}"
        print(f"EXTRACTION_ERROR: {type(e).__name__}: {e} (at {where})")
        raise SystemExit(1)
